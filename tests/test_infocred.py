import shutil
import unittest
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from src.sistemas_externos.infocred import analizar_infocred
from src.sistemas_externos.proceso import ejecutar_revision_completa


class INFOCREDTests(unittest.TestCase):
    @contextmanager
    def _workspace(self):
        root = Path(__file__).resolve().parents[1] / ".tmp-test" / f"infocred-functional-{uuid4().hex}"
        root.mkdir(parents=True)
        try:
            yield root
        finally:
            shutil.rmtree(root)

    def _crear_fuentes(self, root: Path, seccion_codigos: str = "Infocred"):
        matrix_path = root / "matriz.xlsx"
        workbook = Workbook()
        matrix = workbook.active
        matrix.title = "Matriz de Accesos"
        matrix.cell(7, 2, "Analista")
        matrix.cell(7, 16, "300")

        codes = workbook.create_sheet("Codigos")
        codes.cell(1, 1, seccion_codigos)
        codes.cell(1, 2, "Consulta")
        codes.cell(1, 3, "300")
        codes.cell(2, 2, "Rol extra")
        codes.cell(2, 3, "999")
        codes.cell(3, 2, "Rol inactivo")
        codes.cell(3, 3, "999")
        codes.merge_cells("A1:A3")
        workbook.save(matrix_path)

        salar_path = root / "salar.txt"
        salar_rows = [
            ("ANA PERSONA", "Analista", "ana@example.com", "ana"),
            ("BOB PERSONA", "Analista", "bob@example.com", "bob"),
        ]
        lines = []
        for name, cargo, email, login in salar_rows:
            row = [""] * 20
            row[0] = name
            row[13] = "Gerencia"
            row[14] = "Oficina"
            row[16] = cargo
            row[18] = email
            row[19] = login
            lines.append("|".join(row))
        salar_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

        domain_path = root / "dominio.txt"
        domain = [""] * 13
        domain[0] = "CAROL PERSONA"
        domain[2] = "Gerencia"
        domain[3] = "Analista"
        domain[4] = "carol@example.com"
        domain[9] = "Oficina"
        domain[12] = "carol"
        domain_path.write_text("|".join(domain) + "\n", encoding="utf-8")

        infocred_path = root / "infocred.txt"
        txt_rows = []
        for email, role, state in (
            ("ana@example.com", "Consulta", "ACTIVO"),
            ("ana@example.com", "Rol inactivo", "BAJA"),
            ("carol@example.com", "Rol extra", "ACTIVO"),
        ):
            row = [""] * 17
            row[6] = email
            row[8] = role
            row[12] = state
            txt_rows.append("|".join(row))
        infocred_path.write_text("\n".join(txt_rows) + "\n", encoding="utf-8")
        return matrix_path, salar_path, domain_path, infocred_path

    def test_filtra_baja_traduce_roles_y_compara_salar_dominio_y_altas(self):
        with self._workspace() as root:
            matrix, salar, domain, infocred = self._crear_fuentes(root)
            result = analizar_infocred(matrix, salar, infocred, dominio=domain)

        ana_detail = next(item for item in result["detalle"]
                          if item.get("correo") == "ana@example.com")
        self.assertEqual(ana_detail["resultado"], "Conforme")
        self.assertEqual(ana_detail["encontrado"], ["300"])
        self.assertFalse(any(item.get("correo") == "ana@example.com"
                             for item in result["observaciones"]))

        alta = next(item for item in result["observaciones"]
                    if item.get("correo") == "bob@example.com")
        self.assertEqual(alta["categoria"], "Usuarios en SALAR sin registro en sistema")
        self.assertEqual(alta["esperado"], ["300"])

        domain_difference = next(item for item in result["observaciones"]
                                 if item.get("correo") == "carol@example.com")
        self.assertEqual(domain_difference["categoria"], "Usuario Externo (Dominio)")
        self.assertEqual(domain_difference["esperado"], ["300"])
        self.assertEqual(domain_difference["encontrado"], ["999"])

    def test_acepta_el_rotulo_lnfocred_de_la_matriz(self):
        with self._workspace() as root:
            matrix, salar, domain, infocred = self._crear_fuentes(root, "lnfocred")
            result = analizar_infocred(matrix, salar, infocred, dominio=domain)

        ana_detail = next(item for item in result["detalle"]
                          if item.get("correo") == "ana@example.com")
        self.assertEqual(ana_detail["resultado"], "Conforme")
        self.assertEqual(ana_detail["encontrado"], ["300"])

    def test_revision_completa_acepta_solo_infocred(self):
        with self._workspace() as root:
            matrix, salar, domain, infocred = self._crear_fuentes(root)
            output, result = ejecutar_revision_completa(
                matrix, salar, {}, root / "analisis.xlsx", (),
                archivo_dominio=domain, archivo_infocred=infocred)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertIn("Observaciones INFOCRED", workbook.sheetnames)
                self.assertIn("Observaciones Dominio", workbook.sheetnames)
            finally:
                workbook.close()
            self.assertEqual(result["sistemas"], [])
            self.assertGreater(result["infocred"]["usuarios"], 0)


if __name__ == "__main__":
    unittest.main()
