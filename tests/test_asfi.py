import shutil
import unittest
from contextlib import contextmanager
from datetime import date, timedelta
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from src.sistemas_externos.asfi import analizar_asfi


class ASFITemporalesTests(unittest.TestCase):
    @contextmanager
    def _workspace(self):
        root = Path(__file__).resolve().parents[1] / ".tmp-test" / f"asfi-functional-{uuid4().hex}"
        root.mkdir()
        try:
            yield root
        finally:
            shutil.rmtree(root)

    def _crear_fuentes(self, root: Path, matrix_roles, salar_rows, temporal_rows,
                       domain_rows=None):
        matrix_path = root / "matriz.xlsx"
        workbook = Workbook()
        matrix = workbook.active
        matrix.title = "Matriz de Accesos"
        for row_number, (cargo, roles) in enumerate(matrix_roles, start=7):
            matrix.cell(row_number, 2, cargo)
            matrix.cell(row_number, 8, roles)
        codes = workbook.create_sheet("Codigos")
        codes.cell(1, 1, "SISTEMAS ASFI")
        for row_number, role in enumerate(("100", "200", "999"), start=1):
            codes.cell(row_number, 3, role)
            codes.cell(row_number, 4, f"ALIAS{role}")
        workbook.save(matrix_path)

        salar_path = root / "salar.txt"
        salar_lines = []
        for name, cargo, email, login in salar_rows:
            row = [""] * 20
            row[0] = name
            row[13] = "Gerencia"
            row[14] = "Oficina"
            row[16] = cargo
            row[18] = email
            row[19] = login
            salar_lines.append("|".join(row))
        salar_path.write_text("\n".join(salar_lines) + "\n", encoding="utf-8")

        asfi_path = root / "asfi.txt"
        asfi_path.write_text(
            "externo@example.com|Usuario Externo|ALIAS999\n", encoding="utf-8")

        temporal_path = root / "temporales.xlsx"
        temporal_workbook = Workbook()
        temporal_sheet = temporal_workbook.active
        temporal_sheet.append([
            "Usuario", "Nombre", "Tipo", "Cargo Actual", "Cargo Designado",
            "Oficina Actual", "Oficina Designada", "Fecha inicio", "Fecha fin",
        ])
        for login, name, cargo_actual, cargo_designado, fecha_fin, paint_actual, paint_designated in temporal_rows:
            temporal_sheet.append([
                login, name, "Designacion", cargo_actual, cargo_designado,
                "Oficina actual", "Oficina designada", date.today() - timedelta(days=5), fecha_fin,
            ])
            row_number = temporal_sheet.max_row
            green_fill = PatternFill(fill_type="solid", fgColor="FF00FF00")
            if paint_actual:
                temporal_sheet.cell(row_number, 4).fill = green_fill
            if paint_designated:
                temporal_sheet.cell(row_number, 5).fill = green_fill
        temporal_workbook.save(temporal_path)

        domain_path = None
        if domain_rows is not None:
            domain_path = root / "dominio.txt"
            domain_lines = []
            for name, cargo, email, login in domain_rows:
                row = [""] * 13
                row[0] = name
                row[2] = "Gerencia"
                row[3] = cargo
                row[4] = email
                row[9] = "Oficina"
                row[12] = login
                domain_lines.append("|".join(row))
            domain_path.write_text("\n".join(domain_lines) + "\n", encoding="utf-8")
        return matrix_path, salar_path, asfi_path, temporal_path, domain_path

    def _run(self, root, matrix_roles, salar_rows, temporal_rows, domain_rows=None):
        paths = self._crear_fuentes(
            root, matrix_roles, salar_rows, temporal_rows, domain_rows)
        matrix, salar, asfi, temporales, dominio = paths
        return analizar_asfi(matrix, salar, asfi, dominio=dominio,
                             archivo_temporales=temporales)

    def test_alta_salar_usa_cargo_designado_temporal(self):
        with self._workspace() as root:
            result = self._run(root,
                [("Cargo Original", "100"), ("Cargo Designado", "200")],
                [("Ana Persona", "Cargo Original", "ana@example.com", "ana")],
                [("ana", "Ana Persona", "Cargo Original", "Cargo Designado",
                  date.today() + timedelta(days=30), False, True)],
                [("Usuario Externo", "Cargo sin matriz", "externo@example.com", "externo")])

        alta = next(item for item in result["observaciones"]
                    if item.get("correo") == "ana@example.com")
        self.assertEqual(alta["categoria"], "Usuarios en SALAR sin registro en sistema")
        self.assertEqual(alta["cargo"], "Cargo Designado")
        self.assertEqual(alta["esperado"], ["200"])

    def test_alta_salar_vencida_usa_cargo_de_origen(self):
        with self._workspace() as root:
            result = self._run(root,
                [("Cargo Original", "100"), ("Cargo Designado", "200")],
                [("Ana Persona", "Cargo Original", "ana@example.com", "ana")],
                [("ana", "Ana Persona", "Cargo Original", "Cargo Designado",
                  date.today() - timedelta(days=1), False, True)],
                [("Usuario Externo", "Cargo sin matriz", "externo@example.com", "externo")])

        alta = next(item for item in result["observaciones"]
                    if item.get("correo") == "ana@example.com")
        self.assertEqual(alta["cargo"], "Cargo Original")
        self.assertEqual(alta["esperado"], ["100"])

    def test_alta_temporal_con_ambos_cargos_pintados_suma_ambas_expectativas(self):
        with self._workspace() as root:
            result = self._run(root,
                [("Cargo Original", "100"), ("Cargo Designado", "200")],
                [("Ana Persona", "Cargo Original", "ana@example.com", "ana")],
                [("ana", "Ana Persona", "Cargo Original", "Cargo Designado",
                  date.today() + timedelta(days=30), True, True)],
                [("Usuario Externo", "Cargo sin matriz", "externo@example.com", "externo")])

        alta = next(item for item in result["observaciones"]
                    if item.get("correo") == "ana@example.com")
        self.assertEqual(alta["cargo"],
                         "(Actual) Cargo Original\n(Designado) Cargo Designado")
        self.assertEqual(alta["esperado"], ["100", "200"])

    def test_dominio_sin_salar_usa_cargo_designado_temporal(self):
        with self._workspace() as root:
            result = self._run(root,
                [("Cargo Original", "100"), ("Cargo Designado", "200")],
                [("Persona SALAR", "Cargo sin matriz", "salar@example.com", "salar")],
                [("dana", "Dana Persona", "Cargo Original", "Cargo Designado",
                  date.today() + timedelta(days=30), False, True)],
                [("Dana Persona", "Cargo Original", "dana@example.com", "dana")])

        alta = next(item for item in result["observaciones"]
                    if item.get("correo") == "dana@example.com")
        self.assertEqual(alta["categoria"], "Usuario en Dominio sin registro en sistema")
        self.assertEqual(alta["cargo"], "Cargo Designado")
        self.assertEqual(alta["esperado"], ["200"])


if __name__ == "__main__":
    unittest.main()
