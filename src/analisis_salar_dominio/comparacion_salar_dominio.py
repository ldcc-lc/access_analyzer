"""Conciliación de los datos de personas de SALAR y Dominio."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..sistemas_internos.lectores import (
    DOMINIO_CONFIG, SALAR_CONFIG, Persona, PersonasConfig, TemporalesConfig,
    leer_accesos_temporales, leer_personas_txt, leer_salar, normalizar_cargo,
)


DOMAIN_COMPARISON_HEADERS = (
    "CATEGORÍA", "LOGIN", "NOMBRE", "CORREO", "CARGO",
    "OFICINA", "GERENCIA", "VALOR ESPERADO",
    "VALOR ENCONTRADO (DOMINIO)", "FECHA INICIO", "FECHA FIN",
    "CARGO EN MATRIZ",
)
DOMAIN_COMPARISON_WIDTHS = (34, 24, 38, 34, 40, 34, 34, 52, 52, 16, 16, 18)
DOMAIN_DETAIL_HEADERS = (
    "RESULTADO", "COINCIDENCIA POR", "LOGIN SALAR", "LOGIN DOMINIO",
    "VALOR ESPERADO (SALAR)", "VALOR ENCONTRADO (DOMINIO)", "NOMBRE SALAR",
    "NOMBRE DOMINIO", "CORREO SALAR", "CORREO DOMINIO", "CARGO SALAR",
    "CARGO DOMINIO", "OFICINA SALAR", "OFICINA DOMINIO", "GERENCIA SALAR",
    "GERENCIA DOMINIO",
)
DOMAIN_DETAIL_WIDTHS = (30, 23, 22, 22, 42, 42, 34, 34, 30, 30, 38, 38, 30, 30, 30, 30)
HISTORIAL_HEADERS = (
    "RESULTADO", "CARNET", "NOMBRE EN SALAR ANTERIOR", "LOGIN ANTERIOR",
    "CORREO ANTERIOR", "CARGO ANTERIOR", "LOGIN EN DOMINIO",
    "CORREO EN DOMINIO", "CARGO EN DOMINIO", "OFICINA EN DOMINIO",
    "GERENCIA EN DOMINIO", "MÉTODO DE COINCIDENCIA",
)
HISTORIAL_WIDTHS = (36, 18, 36, 22, 34, 40, 22, 34, 40, 34, 34, 24)
LOAD_ERROR_HEADERS = ("ARCHIVO", "NÚMERO DE FILA", "CAMPO AFECTADO", "DESCRIPCIÓN")


def _norm(value) -> str:
    # Reutiliza la normalización de cargos para construir claves de comparación.
    return normalizar_cargo("" if value is None else str(value))


def _normalizar_nombre(value) -> str:
    """Ignore case, accents, punctuation and extra spaces for name matching."""
    text = "" if value is None else str(value)
    text = "".join(char for char in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(char))
    text = "".join(" " if unicodedata.category(char).startswith("P") else char
                   for char in text)
    return " ".join(text.casefold().split())


def _email_key(value) -> str:
    # Canonicaliza el correo ignorando espacios y diferencias entre mayúsculas/minúsculas.
    return " ".join(str(value or "").split()).casefold()


def _carnet_key(value) -> str:
    # Reduce el carnet a caracteres alfanuméricos para equiparar formatos distintos.
    return "".join(char for char in str(value or "").upper() if char.isalnum())


def _analizar_salar_anterior(old_people, current_salar, domain_people):
    """Detect missing unique carnets and check whether the old account remains in Domain."""
    # Usa el carnet como identificador histórico y comprueba si la cuenta sigue en Dominio.
    old_by_id: dict[str, list[Persona]] = defaultdict(list)
    current_by_id: dict[str, list[Persona]] = defaultdict(list)
    rows = []
    for person in old_people:
        key = _carnet_key(person.carnet)
        if key:
            old_by_id[key].append(person)
        else:
            rows.append(("Revisión manual: carnet vacío en SALAR anterior",
                         "", person.nombre, person.login, person.correo,
                         person.cargo, "", "", "", "", "", ""))
    for person in current_salar:
        key = _carnet_key(person.carnet)
        if key:
            current_by_id[key].append(person)

    domain_logins: dict[str, list[Persona]] = defaultdict(list)
    domain_emails: dict[str, list[Persona]] = defaultdict(list)
    for person in domain_people:
        if _norm(person.login):
            domain_logins[_norm(person.login)].append(person)
        if _email_key(person.correo):
            domain_emails[_email_key(person.correo)].append(person)

    for carnet, historical in old_by_id.items():
        if len(historical) != 1:
            for person in historical:
                rows.append(("Revisión manual: carnet duplicado en SALAR anterior",
                             person.carnet, person.nombre, person.login, person.correo,
                             person.cargo, "", "", "", "", "", ""))
            continue
        if carnet in current_by_id:
            if len(current_by_id[carnet]) > 1:
                person = historical[0]
                rows.append(("Revisión manual: carnet duplicado en SALAR actual",
                             person.carnet, person.nombre, person.login, person.correo,
                             person.cargo, "", "", "", "", "", ""))
            continue

        person = historical[0]
        login_key, email_key = _norm(person.login), _email_key(person.correo)
        login_hits = domain_logins.get(login_key, []) if login_key else []
        email_hits = domain_emails.get(email_key, []) if email_key else []
        candidates = {id(hit): hit for hit in login_hits + email_hits}
        unique_methods = [(login_hits[0], "Login")] if len(login_hits) == 1 else []
        if len(email_hits) == 1:
            unique_methods.append((email_hits[0], "Correo"))
        matched = {id(hit): hit for hit, _ in unique_methods}
        if len(candidates) == len(matched) == 1:
            domain = next(iter(matched.values()))
            method = "+".join(dict.fromkeys(label for hit, label in unique_methods
                                             if hit is domain))
            result = "Posible baja SALAR: cuenta aún en Dominio"
            domain_values = (domain.login, domain.correo, domain.cargo,
                             domain.oficina, domain.gerencia)
        elif candidates:
            result, method = "Revisión manual: coincidencia de cuenta ambigua", ""
            domain_values = ("", "", "", "", "")
        else:
            result = "Ausente de SALAR actual; cuenta no encontrada en Dominio"
            method, domain_values = "", ("", "", "", "", "")
        rows.append((result, person.carnet, person.nombre, person.login,
                     person.correo, person.cargo, *domain_values, method))
    return rows


def _domain_config(config: PersonasConfig) -> PersonasConfig:
    # Ajusta los campos de lectura para interpretar el archivo de Dominio.
    """Use the requested default columns while respecting saved settings."""
    return PersonasConfig(
        campos_txt=config.campos_txt,
        campos_nombre=config.campos_nombre,
        campo_cargo=config.campo_cargo,
        campo_login=config.campo_login,
        campo_oficina=config.campo_oficina or 10,
        campo_gerencia=config.campo_gerencia or 3,
        campo_correo=config.campo_correo or 5,
    )


def _persona_values(persona: Persona) -> dict[str, str]:
    # Expone los atributos comunes de Persona para construir filas del reporte.
    return {
        "login": persona.login or "", "nombre": persona.nombre or "",
        "correo": persona.correo or "", "cargo": persona.cargo or "",
        "oficina": persona.oficina or "", "gerencia": persona.gerencia or "",
    }


def _row(category: str, salar: Persona | None, domain: Persona | None,
         field: str = "", expected: str = "", found: str = "",
         fecha_inicio: str = "", fecha_fin: str = "",
         cargo_en_matriz: bool | None = None) -> dict:
    # Construye una observación unificada; para coincidencias conserva los datos SALAR;
    # el login visible, por indicación de la usuaria, siempre es el de Dominio.
    preferred = _persona_values(salar or domain) if (salar or domain) else {}
    if _norm(category) == _norm("Usuario sin Dominio pendiente de creación") and salar:
        # En esta categoría no existe login de Dominio: el identificador que
        # debe conservarse en la comparación es el asignado por SALAR.
        login = salar.login_salar or salar.login
    else:
        login = domain.login if domain else ""
    if field:
        expected_value = f"{field}: {expected}" if expected else ""
        found_value = f"{field}: {found}" if found else ""
    else:
        expected_value = found_value = ""
    return {"values": (
        category, login, preferred.get("nombre", ""),
        preferred.get("correo", ""), preferred.get("cargo", ""),
        preferred.get("oficina", ""), preferred.get("gerencia", ""),
        expected_value, found_value, fecha_inicio, fecha_fin,
        "Sí" if cargo_en_matriz is True else "No" if cargo_en_matriz is False else "",
    ), "field": field, "expected": expected, "found": found,
        "cargo_en_matriz": cargo_en_matriz}


def _match_people(salar_people: list[Persona], domain_people: list[Persona]):
    # Concilia identidades por claves fuertes antes de intentar coincidencias por nombre.
    """Match by login, email, then unique name if SALAR has no email."""
    salar_login: dict[str, list[int]] = defaultdict(list)
    domain_login: dict[str, list[int]] = defaultdict(list)
    for i, person in enumerate(salar_people):
        key = _norm(person.login)
        if key:
            salar_login[key].append(i)
    for i, person in enumerate(domain_people):
        key = _norm(person.login)
        if key:
            domain_login[key].append(i)

    pairs: list[tuple[int, int]] = []
    used_salar: set[int] = set()
    used_domain: set[int] = set()
    methods: dict[tuple[int, int], str] = {}

    def pair_unique(left_map, right_map, method):
        for key in left_map.keys() & right_map.keys():
            left = [i for i in left_map[key] if i not in used_salar]
            right = [i for i in right_map[key] if i not in used_domain]
            if len(left) == len(right) == 1:
                pairs.append((left[0], right[0]))
                methods[(left[0], right[0])] = method
                used_salar.add(left[0])
                used_domain.add(right[0])

    pair_unique(salar_login, domain_login, "Login")

    salar_email: dict[str, list[int]] = defaultdict(list)
    domain_email: dict[str, list[int]] = defaultdict(list)
    for i, person in enumerate(salar_people):
        key = _email_key(person.correo)
        if i not in used_salar and key:
            salar_email[key].append(i)
    for i, person in enumerate(domain_people):
        key = _email_key(person.correo)
        if i not in used_domain and key:
            domain_email[key].append(i)
    pair_unique(salar_email, domain_email, "Correo")

    # Si SALAR no tiene correo, un nombre completo normalizado y único puede
    # reconocer la cuenta de Dominio aunque el login todavía sea distinto.
    salar_names: dict[str, list[int]] = defaultdict(list)
    domain_names: dict[str, list[int]] = defaultdict(list)
    for i, person in enumerate(salar_people):
        key = _normalizar_nombre(person.nombre)
        if key:
            salar_names[key].append(i)
    for i, person in enumerate(domain_people):
        key = _normalizar_nombre(person.nombre)
        if key:
            domain_names[key].append(i)
    for key in salar_names.keys() & domain_names.keys():
        left, right = salar_names[key], domain_names[key]
        if len(left) == len(right) == 1:
            si, di = left[0], right[0]
            if (si not in used_salar and di not in used_domain
                    and not _email_key(salar_people[si].correo)):
                pairs.append((si, di))
                methods[(si, di)] = "Nombre (correo SALAR vacío)"
                used_salar.add(si)
                used_domain.add(di)
    return pairs, used_salar, used_domain, methods


def _detail_row(result, match_method, field, salar, domain, expected="", found=""):
    # Crea una fila con el método de cruce y los valores de ambas fuentes.
    sv = _persona_values(salar) if salar else {}
    dv = _persona_values(domain) if domain else {}
    expected_value = f"{field}: {expected}" if field and expected else ""
    found_value = f"{field}: {found}" if field and found else ""
    return (result, match_method,
            sv.get("login", ""), dv.get("login", ""), expected_value, found_value,
            sv.get("nombre", ""), dv.get("nombre", ""),
            sv.get("correo", ""), dv.get("correo", ""),
            sv.get("cargo", ""), dv.get("cargo", ""),
            sv.get("oficina", ""), dv.get("oficina", ""),
            sv.get("gerencia", ""), dv.get("gerencia", ""))


def _regular_observations(salar_people: list[Persona], domain_people: list[Persona],
                          cargos_matriz: set[str] | None = None):
    # Clasifica coincidencias, diferencias y personas presentes solo en una fuente.
    pairs, used_salar, used_domain, methods = _match_people(salar_people, domain_people)
    rows, details = [], []
    domain_by_name: dict[str, list[int]] = defaultdict(list)
    salar_by_name: dict[str, list[int]] = defaultdict(list)
    for i, person in enumerate(domain_people):
        key = _normalizar_nombre(person.nombre)
        if key:
            domain_by_name[key].append(i)
    for i, person in enumerate(salar_people):
        key = _normalizar_nombre(person.nombre)
        if key:
            salar_by_name[key].append(i)
    fields = (
        ("login", "Login", "login"),
        ("nombre", "Diferencia por nombre"),
        ("cargo", "Diferencia por cargo"),
        ("oficina", "Diferencia por oficina/gerencia"),
        ("gerencia", "Diferencia por oficina/gerencia"),
        ("correo", "Correo", "correo"),
    )
    for si, di in pairs:
        salar, domain = salar_people[si], domain_people[di]
        sv, dv = _persona_values(salar), _persona_values(domain)
        method = methods[(si, di)]
        for item in fields:
            field, category = item[:2]
            normalize = _normalizar_nombre if field == "nombre" else _norm
            if field in {"correo", "login"}:
                normalize = _email_key if field == "correo" else _norm
            differs = normalize(sv[field]) != normalize(dv[field])
            status = category if differs else "Coincide"
            details.append(_detail_row(status, method, field.capitalize(), salar,
                                       domain, sv[field], dv[field]))
            if differs and field in {"nombre", "cargo", "oficina", "gerencia"}:
                rows.append(_row(category, salar, domain, field.capitalize(),
                                 sv[field], dv[field]))

    # SALAR sin correo y sin nombre en Dominio: cuenta pendiente de creación.
    possible_domain_by_name: set[int] = set()
    for i, salar in enumerate(salar_people):
        if i in used_salar:
            continue
        key = _normalizar_nombre(salar.nombre)
        name_candidates = domain_by_name.get(key, [])
        if not _email_key(salar.correo) and key and not name_candidates:
            rows.append(_row("Usuario sin Dominio pendiente de creación", salar, None))
            details.append(_detail_row("Usuario sin Dominio pendiente de creación",
                                       "", "", salar, None))
        elif name_candidates:
            possible_domain_by_name.update(name_candidates)
            details.append(_detail_row("Revisión manual: nombre ambiguo", "Nombre",
                                       "Nombre", salar, None, salar.nombre, " | ".join(
                                           domain_people[j].nombre for j in name_candidates)))
        else:
            details.append(_detail_row("Solo en SALAR", "", "", salar, None))

    # Solo se clasifican como externos los registros de Dominio sin persona
    # SALAR asociada por login/correo/nombre. Nombres candidatos ambiguos quedan
    # en el detalle para evitar falsos positivos.
    for i, domain in enumerate(domain_people):
        if i not in used_domain:
            if i in possible_domain_by_name:
                details.append(_detail_row("Revisión manual: nombre ambiguo", "Nombre",
                                           "Nombre", None, domain, "", domain.nombre))
            else:
                en_matriz = (_norm(domain.cargo) in cargos_matriz
                             if cargos_matriz is not None else None)
                rows.append(_row("Usuario externo", None, domain,
                                 cargo_en_matriz=en_matriz))
                details.append(_detail_row("Usuario externo", "", "", None, domain))
    return rows, details, pairs, methods


def _temporary_observations(salar_people: list[Persona], domain_people: list[Persona],
                            path: str | Path | None,
                            config: TemporalesConfig) -> tuple[list[dict], set[tuple[str, str]]]:
    # Contrasta asignaciones temporales con Dominio y conserva su vigencia en los resultados.
    if not path:
        return [], set()
    temporales = leer_accesos_temporales(path, config)
    pairs, _, _, _ = _match_people(salar_people, domain_people)
    salar_to_domain = {si: di for si, di in pairs}
    domain_to_salar = {di: si for si, di in pairs}
    salar_by_login: dict[str, list[int]] = defaultdict(list)
    domain_by_login: dict[str, list[int]] = defaultdict(list)
    for i, person in enumerate(salar_people):
        if _norm(person.login):
            salar_by_login[_norm(person.login)].append(i)
    for i, person in enumerate(domain_people):
        if _norm(person.login):
            domain_by_login[_norm(person.login)].append(i)

    today = date.today()
    rows = []
    temporal_fields: set[tuple[str, str]] = set()
    for temp in temporales:
        key = _norm(temp["login"])
        salar_hits = salar_by_login.get(key, [])
        domain_hits = domain_by_login.get(key, [])
        salar = salar_people[salar_hits[0]] if len(salar_hits) == 1 else None
        domain = domain_people[domain_hits[0]] if len(domain_hits) == 1 else None
        # El login temporal podría ser el de SALAR; en ese caso se obtiene el
        # registro de Dominio mediante la pareja ya conciliada.
        if salar is not None and domain is None:
            domain_index = salar_to_domain.get(salar_hits[0])
            if domain_index is not None:
                domain = domain_people[domain_index]
        if domain is not None and salar is None:
            salar_index = domain_to_salar.get(domain_hits[0])
            if salar_index is not None:
                salar = salar_people[salar_index]
        if domain is None:
            continue

        start = date.fromisoformat(temp["fecha_inicio"]) if temp.get("fecha_inicio") else None
        end = date.fromisoformat(temp["fecha_fin"]) if temp.get("fecha_fin") else None
        if end and today > end:
            category = "Fin Acceso Temporal"
            expected_cargo = temp.get("cargo_actual", "")
            expected_office = temp.get("oficina_actual", "")
        else:
            # Aún no comenzó la asignación temporal: no se evalúa como vigente.
            if start and today < start:
                continue
            category = "Acceso Temporal"
            both_cargos_painted = (temp.get("cargo_actual_pintado") and
                                   temp.get("cargo_designado_pintado"))
            designated_painted = temp.get("cargo_designado_pintado")
            if not designated_painted:
                continue
            if both_cargos_painted:
                expected_cargo = temp.get("cargo_actual", "")
                expected_office = temp.get("oficina_actual", "")
            else:
                expected_cargo = temp.get("cargo_designado", "")
                expected_office = temp.get("oficina_designada", "")

        # La temporal es la fuente de expectativa para los campos evaluados.
        domain_key = _norm(domain.login)
        for field, expected, found in (
            ("Cargo", expected_cargo, domain.cargo),
            ("Oficina", expected_office, domain.oficina),
        ):
            if not expected:
                continue
            temporal_fields.add((domain_key, field))
            coincide = _norm(expected) == _norm(found)
            estado = "conforme" if coincide else "diferencia"
            prefijo = "Fin Acceso Temporal" if end and today > end else "Acceso Temporal"
            temporal_category = f"{prefijo} {estado.capitalize()}"
            rows.append(_row(temporal_category, salar, domain, field, expected, found,
                             temp.get("fecha_inicio", ""), temp.get("fecha_fin", "")))
    return rows, temporal_fields


def comparar_salar_dominio(salar_file: str | Path, domain_file: str | Path,
                           config_salar: PersonasConfig = SALAR_CONFIG,
                           config_dominio: PersonasConfig = DOMINIO_CONFIG,
                           archivo_temporales: str | Path | None = None,
                           config_temporales: TemporalesConfig = TemporalesConfig(),
                           archivo_salar_anterior: str | Path | None = None,
                           cargos_matriz: set[str] | None = None
) -> dict:
    """Compara por login/correo y devuelve observaciones, detalle y errores."""
    # Lee las fuentes y coordina las observaciones ordinarias, históricas y temporales.
    salar_people, salar_errors = leer_salar(salar_file, config_salar)
    domain_people, domain_errors = leer_personas_txt(domain_file, _domain_config(config_dominio))
    rows, details, _, methods = _regular_observations(
        salar_people, domain_people, cargos_matriz)
    temporal_rows, temporal_fields = _temporary_observations(
        salar_people, domain_people, archivo_temporales, config_temporales)
    # No repetir en las observaciones ordinarias los campos que ya evalúa una
    # asignación temporal vigente o vencida.
    rows = [row for row in rows if
            (_norm(row["values"][1]), row.get("field", "")) not in temporal_fields]
    rows.extend(temporal_rows)
    for row in temporal_rows:
        values = row["values"]
        domain_login = values[1]
        domain = next((person for person in domain_people if person.login == domain_login), None)
        salar = None
        if domain:
            pair = next(((si, di) for (si, di) in methods if domain_people[di] is domain), None)
            if pair:
                salar = salar_people[pair[0]]
        details.append(_detail_row(values[0], "Temporal", row["field"], salar, domain,
                                   row["expected"], row["found"]))
    errors = salar_errors + domain_errors
    resultado = {"observaciones": rows, "detalle": details, "errores": errors}
    if archivo_salar_anterior:
        historical_people, historial_errors = leer_salar(
            archivo_salar_anterior, config_salar)
        errors.extend(historial_errors)
        resultado["bajas_historicas"] = _analizar_salar_anterior(
            historical_people, salar_people, domain_people)
    return resultado


def agregar_hoja_comparacion(path: str | Path, resultado: dict) -> None:
    # Agrega hojas de conciliación al Excel generado por el análisis principal.
    """Add comparison observations, detail, and source load errors to workbook."""
    path = Path(path)
    workbook = load_workbook(path, rich_text=True)
    try:
        title = "Observaciones Dominio"
        temporal_title = "Dominio Usuarios Temporales"
        for existing_title in (title, temporal_title, "Comparación SALAR-Dominio",
                               "Detalle SALAR-Dominio", "Errores SALAR-Dominio",
                               "Historial SALAR-Dominio"):
            if existing_title in workbook.sheetnames:
                del workbook[existing_title]
        sheet = workbook.create_sheet(title)
        sheet.append(DOMAIN_COMPARISON_HEADERS)
        temporal_rows = [row for row in resultado["observaciones"]
                         if str(row["values"][0]).startswith(
                             ("Acceso Temporal", "Fin Acceso Temporal"))]
        ordinary_rows = [row for row in resultado["observaciones"]
                         if row not in temporal_rows]
        for row in ordinary_rows:
            sheet.append(row["values"])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
        for column, width in enumerate(DOMAIN_COMPARISON_WIDTHS, start=1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
        sheet.row_dimensions[1].height = 34
        category_fills = {
            "Usuario externo": "D9EAD3",
            "Usuario sin Dominio pendiente de creación": "A9D18E",
            "Diferencia por oficina/gerencia": "FCE4D6",
            "Diferencia por nombre": "FFF2CC",
            "Diferencia por cargo": "DDEBF7",
            "Acceso Temporal Conforme": "A9D18E",
            "Acceso Temporal Diferencia": "FCE4D6",
            "Fin Acceso Temporal Conforme": "F4CCCC",
            "Fin Acceso Temporal Diferencia": "F4CCCC",
        }
        for row_number in range(2, sheet.max_row + 1):
            category = sheet.cell(row_number, 1).value
            if category in category_fills:
                sheet.cell(row_number, 1).fill = PatternFill(
                    "solid", fgColor=category_fills[category])
            max_lines = 1
            for column, width in enumerate(DOMAIN_COMPARISON_WIDTHS, start=1):
                cell = sheet.cell(row_number, column)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                max_lines = max(max_lines, sum(
                    max(1, (len(line) + width - 1) // width)
                    for line in str(cell.value or "").splitlines()
                ))
            sheet.row_dimensions[row_number].height = max(22, max_lines * 15 + 6)

        temporal_sheet = workbook.create_sheet(temporal_title)
        temporal_sheet.append(DOMAIN_COMPARISON_HEADERS)
        for row in temporal_rows:
            temporal_sheet.append(row["values"])
        _format_sheet(temporal_sheet, DOMAIN_COMPARISON_WIDTHS)
        for row_number, row in enumerate(temporal_rows, start=2):
            color = category_fills.get(row["values"][0])
            if color:
                temporal_sheet.cell(row_number, 1).fill = PatternFill(
                    "solid", fgColor=color)

        detail = workbook.create_sheet("Detalle SALAR-Dominio")
        detail.append(DOMAIN_DETAIL_HEADERS)
        for row in resultado["detalle"]:
            detail.append(row)
        _format_sheet(detail, DOMAIN_DETAIL_WIDTHS)

        errors = workbook.create_sheet("Errores SALAR-Dominio")
        errors.append(LOAD_ERROR_HEADERS)
        for item in resultado["errores"]:
            errors.append(tuple(item.get(key, "") for key in
                                ("archivo", "fila", "campo", "descripcion")))
        _format_sheet(errors, (32, 16, 28, 60))

        if "bajas_historicas" in resultado:
            history = workbook.create_sheet("Historial SALAR-Dominio")
            history.append(HISTORIAL_HEADERS)
            for row in resultado["bajas_historicas"]:
                history.append(row)
            _format_sheet(history, HISTORIAL_WIDTHS)
            for row_number in range(2, history.max_row + 1):
                result = str(history.cell(row_number, 1).value or "")
                color = ("F4CCCC" if result.startswith("Posible baja") else
                         "FFF2CC" if result.startswith("Revisión manual") else None)
                if color:
                    history.cell(row_number, 1).fill = PatternFill("solid", fgColor=color)
        workbook.save(path)
    finally:
        workbook.close()


def _format_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    sheet.row_dimensions[1].height = 34
    for row_number in range(2, sheet.max_row + 1):
        max_lines = 1
        for column, width in enumerate(widths, start=1):
            cell = sheet.cell(row_number, column)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            max_lines = max(max_lines, sum(max(1, (len(line) + width - 1) // width)
                                           for line in str(cell.value or "").splitlines()))
        sheet.row_dimensions[row_number].height = max(22, max_lines * 15 + 6)
