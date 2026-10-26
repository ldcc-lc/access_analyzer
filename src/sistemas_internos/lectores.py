"""Lectura y validación de la matriz Excel y los TXT delimitados por |."""
from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path



class ErrorCarga(ValueError):
    """Error de estructura que impide una comparación confiable."""


# Estas configuraciones describen la estructura de cada archivo de entrada.
# Indica que hoja y columna de cargo se deben leer en la matriz.
@dataclass(frozen=True)
class MatrizConfig:
    hoja: str = "Matriz de Accesos"
    columna_cargo: str = "B"


# Configura cuantas columnas tiene el TXT de personas y donde esta cada dato.
@dataclass(frozen=True)
class PersonasConfig:
    campos_txt: int
    campos_nombre: tuple[int, ...]
    campo_cargo: int
    campo_login: int
    campo_oficina: int | None = None
    campo_gerencia: int | None = None
    campo_correo: int | None = None
    campo_carnet: int | None = None


# Configura encabezado y columnas del archivo de accesos temporales.
@dataclass(frozen=True)
class TemporalesConfig:
    fila_encabezado: int = 1
    campo_login: int = 1
    campo_nombre: int = 2
    campo_cargo_actual: int = 4
    campo_cargo_designado: int = 5
    campo_oficina_actual: int = 6
    campo_oficina_designada: int = 7
    campo_fecha_inicio: int = 8
    campo_fecha_fin: int = 9


DOMINIO_CONFIG = PersonasConfig(
    13, (1,), 4, 13, campo_oficina=10, campo_gerencia=3, campo_correo=5)
SALAR_CONFIG = PersonasConfig(20, (1, 2, 3, 4, 5), 17, 20, 15, 14, 19,
                              campo_carnet=7)


# Modelos de datos comunes: mantienen los registros independientes del formato TXT/Excel.
# Representa una persona con identidad y datos organizacionales ya extraidos.
@dataclass
class Persona:
    login: str
    cargo: str
    nombre: str = ""
    fila: int = 0
    oficina: str = ""
    gerencia: str = ""
    correo: str = ""
    login_salar: str = ""
    sin_coincidencia_dominio: bool = False
    carnet: str = ""


# Representa una cuenta encontrada en el TXT de un sistema y sus accesos.
@dataclass
class CuentaSistema:
    login: str
    roles_texto: str
    fila: int = 0
    organigrama_texto: str = ""
    nombre: str = ""
    cargo: str = ""
    oficina: str = ""
    gerencia: str = ""
    estado: str = ""
    login_respaldo: str = ""


# Define como leer, cruzar y comparar un sistema con la matriz.
@dataclass(frozen=True)
class SistemaConfig:
    nombre: str
    archivo: str
    columna_matriz: str
    encabezado_matriz: str
    campo_login: int
    campo_roles: int
    campos_txt: int
    separador_roles: str = ","
    prefijo_login: str = ""
    sufijo_login: str = ""
    null_sin_roles: bool = False
    columna_matriz_organigrama: str | None = None
    encabezado_matriz_organigrama: str = ""
    campo_organigrama: int | None = None
    separador_organigrama_matriz: str = "-"
    separador_organigrama_txt: str = ","
    campo_estado: int | None = None
    estado_activo: str = ""
    estado_inactivo: str = ""
    agrupar_roles_por_login: bool = False
    normalizar_codigos_rol: bool = False
    campo_nombre: int | None = None
    campo_cargo: int | None = None
    campo_oficina: int | None = None
    campo_gerencia: int | None = None
    clave_cruce: str = "login"
    validar_solo_existencia: bool = False
    campo_login_respaldo: int | None = None
    prefijo_login_respaldo: str = ""


# Guarda los accesos que la matriz asigna a un cargo.
@dataclass
class ReglaCargo:
    cargo: str
    roles_texto: str
    fila: int = 0
    organigrama_texto: str = ""


# Normaliza cargos para comparar textos equivalentes sin diferencias de formato.
def normalizar_cargo(valor: str) -> str:
    # Elimina diferencias de espacios, mayúsculas, acentos y variantes conocidas del cargo.
    texto = " ".join((valor or "").split()).casefold()
    texto = "".join(c for c in unicodedata.normalize("NFKD", texto)
                    if not unicodedata.combining(c))
    texto = re.sub(r"[’‘'ʼ]", "", texto)
    texto = re.sub(r"\bsub\s+gerente\b", "subgerente", texto)
    texto = re.sub(r"\bntic\s+s\b", "ntics", texto)
    return texto


# Uniforma un login para usarlo como clave de identidad.
def normalizar_login(valor: str) -> str:
    # Normaliza la identidad sin modificar el valor original guardado en el registro.
    return (valor or "").strip().upper()


# Aplica los ajustes propios del sistema antes de cruzar logins.
def normalizar_login_sistema(valor: str, config: SistemaConfig) -> str:
    # Ajusta prefijo/sufijo específicos antes de comparar logins de un sistema.
    login = normalizar_login(valor)
    prefix = (config.prefijo_login or "").upper()
    suffix = (config.sufijo_login or "").upper()
    if prefix and login.startswith(prefix):
        login = login[len(prefix):]
    if suffix and login.endswith(suffix):
        login = login[:-len(suffix)]
    return login


# Lee el TXT, prueba codificaciones conocidas y valida cantidad de columnas.
def _leer_txt(path: Path, campos: int) -> list[list[str]]:
    # Abre TXT con codificaciones habituales, valida columnas y conserva el número de fila.
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ErrorCarga(f"No se pudo abrir {path}: {exc}") from exc
    texto = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if texto is None:
        raise ErrorCarga(f"{path}: codificación no reconocida; no se leyó el archivo completo.")
    filas = []
    for numero, row in enumerate(csv.reader(texto.splitlines(), delimiter="|"), 1):
        if not row or (len(row) == 1 and not row[0].strip()):
            continue
        if len(row) != campos:
            raise ErrorCarga(f"{path}, fila {numero}: se esperaban {campos} campos y llegaron {len(row)}.")
        filas.append((numero, row))
    if not filas:
        raise ErrorCarga(f"{path}: el archivo no contiene registros.")
    return filas


# Extrae personas usando las posiciones declaradas en PersonasConfig.
def leer_personas_txt(path: str | Path, config: PersonasConfig = DOMINIO_CONFIG) -> tuple[list[Persona], list[dict]]:
    # Convierte cada fila de personas en Persona usando posiciones configurables (base 1).
    path = Path(path)
    posiciones = (*config.campos_nombre, config.campo_cargo, config.campo_login,
                  *(p for p in (config.campo_oficina, config.campo_gerencia,
                                 config.campo_correo, config.campo_carnet) if p))
    if config.campos_txt < 1 or any(posicion < 1 or posicion > config.campos_txt
                                    for posicion in posiciones):
        raise ErrorCarga(f"Configuración de personas inválida para {path.name}: los campos deben estar entre 1 y {config.campos_txt}.")
    filas = _leer_txt(path, config.campos_txt)
    personas, errores = [], []
    for numero, row in filas:
        login, cargo = row[config.campo_login - 1], row[config.campo_cargo - 1]
        nombre = " ".join(row[posicion - 1].strip() for posicion in config.campos_nombre
                           if row[posicion - 1].strip())
        if not login.strip():
            errores.append({"archivo": path.name, "fila": numero,
                "campo": f"{config.campo_login} ({config.clave_cruce})",
                "descripcion": f"{config.clave_cruce.capitalize()} vacío."})
        if not cargo.strip():
            errores.append({"archivo": path.name, "fila": numero,
                "campo": f"{config.campo_cargo} (cargo)", "descripcion": "Cargo vacío."})
        oficina = row[config.campo_oficina - 1] if config.campo_oficina else ""
        gerencia = row[config.campo_gerencia - 1] if config.campo_gerencia else ""
        correo = row[config.campo_correo - 1].strip() if config.campo_correo else ""
        # En las exportaciones de personas, "Null" representa un correo ausente.
        # No debe usarse como clave compartida ni producir falsos duplicados.
        if correo.casefold() == "null":
            correo = ""
        carnet = row[config.campo_carnet - 1] if config.campo_carnet else ""
        personas.append(Persona(login, cargo, nombre, numero, oficina, gerencia,
                                correo, carnet=str(carnet or "").strip()))
    return personas, errores


def leer_salar(path: str | Path, config: PersonasConfig = SALAR_CONFIG) -> tuple[list[Persona], list[dict]]:
    """Lee personas desde SALAR usando los campos indicados en su configuración."""
    return leer_personas_txt(path, config)


def leer_sistema_txt(path: str | Path, config: SistemaConfig) -> tuple[list[CuentaSistema], list[dict]]:
    """Lee un sistema desde sus campos configurados, sin interpretar otras columnas."""
    path = Path(path)
    filas = _leer_txt(path, config.campos_txt)
    cuentas, errores = [], []
    agrupadas: dict[str, tuple[str, list[str], int, str]] = {}
    # Extrae solo las columnas declaradas por el sistema y aplica sus filtros de estado.
    for numero, row in filas:
        login = row[config.campo_login - 1]
        login_respaldo = (row[config.campo_login_respaldo - 1]
                          if config.campo_login_respaldo else "")
        roles = row[config.campo_roles - 1]
        estado = row[config.campo_estado - 1].strip() if config.campo_estado else ""
        if config.campo_estado and config.estado_inactivo and estado.casefold() == config.estado_inactivo.strip().casefold():
            continue
        if (config.campo_estado and config.estado_activo and not config.estado_inactivo
                and not config.validar_solo_existencia
                and estado.casefold() != config.estado_activo.strip().casefold()):
            continue
        organigrama = row[config.campo_organigrama - 1] if config.campo_organigrama else ""
        nombre = row[config.campo_nombre - 1] if config.campo_nombre else ""
        cargo = row[config.campo_cargo - 1] if config.campo_cargo else ""
        oficina = row[config.campo_oficina - 1] if config.campo_oficina else ""
        gerencia = row[config.campo_gerencia - 1] if config.campo_gerencia else ""
        if not login.strip() and not login_respaldo.strip():
            errores.append({"archivo": path.name, "fila": numero,
                "campo": f"{config.campo_login} ({config.clave_cruce})",
                "descripcion": f"{config.clave_cruce.capitalize()} vacío."})
        if config.normalizar_codigos_rol:
            normalizado = _normalizar_codigos_rol(roles)
            if normalizado is None:
                errores.append({"archivo": path.name, "fila": numero,
                    "campo": f"{config.campo_roles} (rol)",
                    "descripcion": f"Código de rol con formato de importe no reconocido: {roles!r}."})
            else:
                roles = normalizado
        # Algunos TXT repiten una cuenta por rol; en ese caso se consolida por login.
        if config.agrupar_roles_por_login:
            key = normalizar_login_sistema(login, config)
            if key in agrupadas:
                original_login, roles_acumulados, fila_inicial, nombre = agrupadas[key]
                for rol in roles.split(","):
                    if rol and rol not in roles_acumulados:
                        roles_acumulados.append(rol)
                agrupadas[key] = (original_login, roles_acumulados, fila_inicial, nombre)
            else:
                agrupadas[key] = (login, [rol for rol in roles.split(",") if rol], numero, "")
        else:
            cuentas.append(CuentaSistema(login, roles, numero, organigrama,
                                       nombre, cargo, oficina, gerencia, estado,
                                       login_respaldo))
    if config.agrupar_roles_por_login:
        cuentas = [CuentaSistema(login, ",".join(roles), fila)
                   for login, roles, fila, _ in agrupadas.values()]
    return cuentas, errores


def _normalizar_codigo_rol(valor: str) -> str | None:
    """Convierte importes/códigos como 3.000,00 en el código 3000."""
    texto = (valor or "").strip()
    if re.fullmatch(r"\d+", texto):
        return str(int(texto))
    match = re.fullmatch(r"(\d{1,3}(?:\.\d{3})*|\d+),00", texto)
    if not match:
        return None
    return str(int(match.group(1).replace(".", "")))


def _normalizar_codigos_rol(valor: str) -> str | None:
    """Normaliza uno o varios códigos con formato de importe separados por coma o guion."""
    texto = (valor or "").strip()
    if not texto:
        return ""
    if texto.casefold() == "null" or texto == "-":
        return texto
    patron_importe = r"(?:\d{1,3}(?:\.\d{3})+|\d+),00"
    if re.fullmatch(rf"{patron_importe}(?:[, -]+{patron_importe})*", texto):
        partes = re.findall(patron_importe, texto)
    elif re.fullmatch(r"\d+(?:[, -]+\d+)*", texto):
        partes = re.split(r"[, -]+", texto)
    elif re.fullmatch(patron_importe, texto):
        partes = [texto]
    else:
        return None
    normalizados = [_normalizar_codigo_rol(parte) for parte in partes]
    if any(parte is None for parte in normalizados):
        return None
    return ",".join(dict.fromkeys(normalizados))


# Lee una vez la matriz y devuelve las reglas de los sistemas seleccionados.
def leer_matrices(path: str | Path, configuraciones: tuple[SistemaConfig, ...],
                  config_matriz: MatrizConfig = MatrizConfig()) -> dict[str, list[ReglaCargo]]:
    # Lee una sola hoja de matriz y extrae las reglas para todos los sistemas seleccionados.
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    path = Path(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorCarga(f"No se pudo leer el Excel {path}: {exc}") from exc
    hoja = next((ws for ws in wb.worksheets if ws.title.strip().casefold() == config_matriz.hoja.strip().casefold()), None)
    if hoja is None:
        raise ErrorCarga(f"{path}: no se encontró la hoja {config_matriz.hoja!r}.")
    max_col = max([column_index_from_string(config_matriz.columna_cargo), *(
        column_index_from_string(sistema.columna_matriz_organigrama or sistema.columna_matriz)
        for sistema in configuraciones)])
    filas = list(hoja.iter_rows(min_row=1, min_col=1, max_col=max_col, values_only=True))
    resultado = {}
    # Localiza cada encabezado en su columna y toma las filas de cargo que siguen.
    for sistema_config in configuraciones:
        idx = column_index_from_string(sistema_config.columna_matriz) - 1
        encabezado = next((i for i, values in enumerate(filas)
                           if len(values) > idx and str(values[idx] or "").strip().casefold()
                           == sistema_config.encabezado_matriz.strip().casefold()), None)
        if encabezado is None:
            wb.close()
            raise ErrorCarga(f"{path}: no se encontró el encabezado {sistema_config.encabezado_matriz!r} "
                             f"en la columna {sistema_config.columna_matriz} ({sistema_config.nombre}).")
        idx_org = (column_index_from_string(sistema_config.columna_matriz_organigrama) - 1
                   if sistema_config.columna_matriz_organigrama else None)
        if idx_org is not None:
            org_header = next((i for i, values in enumerate(filas)
                               if len(values) > idx_org and str(values[idx_org] or "").strip().casefold()
                               == sistema_config.encabezado_matriz_organigrama.strip().casefold()), None)
            if org_header != encabezado:
                wb.close()
                raise ErrorCarga(f"{path}: el encabezado {sistema_config.encabezado_matriz_organigrama!r} "
                                 f"en la columna {sistema_config.columna_matriz_organigrama} no está en la fila "
                                 f"del encabezado de {sistema_config.nombre}.")
        reglas = []
        for numero, valores in enumerate(filas[encabezado + 1:], encabezado + 2):
            cargo_val, acceso = valores[column_index_from_string(config_matriz.columna_cargo) - 1], valores[idx]
            organigrama = valores[idx_org] if idx_org is not None else None
            cargo = "" if cargo_val is None else str(cargo_val).strip()
            acceso_texto = "" if acceso is None else str(acceso).strip()
            organigrama_texto = "" if organigrama is None else str(organigrama).strip()
            if not cargo or normalizar_cargo(cargo) in {"cargo", "cargos"}:
                continue
            reglas.append(ReglaCargo(cargo, acceso_texto, numero, organigrama_texto))
        if not reglas:
            wb.close()
            raise ErrorCarga(f"{path}: no se encontraron cargos en las columnas {config_matriz.columna_cargo} y {sistema_config.columna_matriz}.")
        resultado[sistema_config.nombre] = reglas
    wb.close()
    return resultado


def leer_accesos_temporales(path: str | Path,
                            config: TemporalesConfig = TemporalesConfig()) -> list[dict[str, str]]:
    """Lee login, asignaciones y vigencia de las columnas configuradas."""
    from openpyxl import load_workbook
    path = Path(path)
    try:
        # En esta hoja los estilos de las celdas de cargo codifican la regla
        # de roles esperados, por eso se requiere acceso a los estilos.
        wb = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise ErrorCarga(f"No se pudo leer el Excel de accesos temporales {path}: {exc}") from exc
    hoja = wb.active
    resultado = []
    posiciones = (config.campo_login, config.campo_nombre, config.campo_cargo_actual,
                  config.campo_cargo_designado, config.campo_oficina_actual,
                  config.campo_oficina_designada, config.campo_fecha_inicio,
                  config.campo_fecha_fin)
    if config.fila_encabezado < 0 or any(posicion < 1 for posicion in posiciones):
        wb.close()
        raise ErrorCarga(f"Configuración inválida para el archivo de accesos temporales {path}.")
    # Cada fila válida queda lista para asociarse a una cuenta durante la comparación.
    for numero, row in enumerate(hoja.iter_rows(min_row=config.fila_encabezado + 1,
            min_col=1, max_col=max(posiciones)), config.fila_encabezado + 1):
        valores = [celda.value for celda in row]
        login = valores[config.campo_login - 1]
        nombre = valores[config.campo_nombre - 1]
        cargo_actual = valores[config.campo_cargo_actual - 1]
        cargo_designado = valores[config.campo_cargo_designado - 1]
        oficina_actual = valores[config.campo_oficina_actual - 1]
        oficina_designada = valores[config.campo_oficina_designada - 1]
        fecha_inicio = valores[config.campo_fecha_inicio - 1]
        fecha_fin = valores[config.campo_fecha_fin - 1]
        login = "" if login is None else str(login).strip()
        cargo_actual = "" if cargo_actual is None else str(cargo_actual).strip()
        cargo_designado = "" if cargo_designado is None else str(cargo_designado).strip()
        oficina_actual = "" if oficina_actual is None else str(oficina_actual).strip()
        oficina_designada = "" if oficina_designada is None else str(oficina_designada).strip()
        if not login and not cargo_actual and not cargo_designado:
            continue
        if not login or not cargo_designado:
            wb.close()
            raise ErrorCarga(f"{path}, fila {numero}: falta Usuario o Cargo Designado.")
        resultado.append({"login": login,
                          "nombre": "" if nombre is None else str(nombre).strip(),
                          "cargo_actual": cargo_actual,
                          "cargo_designado": cargo_designado,
                          "cargo_actual_pintado": _celda_pintada(row[config.campo_cargo_actual - 1]),
                          "cargo_designado_pintado": _celda_pintada(row[config.campo_cargo_designado - 1]),
                          "oficina_actual": oficina_actual,
                          "oficina_designada": oficina_designada,
                          "fecha_inicio": _fecha_temporal(fecha_inicio, wb.epoch, path, numero,
                                                           config.campo_fecha_inicio),
                          "fecha_fin": _fecha_temporal(fecha_fin, wb.epoch, path, numero,
                                                        config.campo_fecha_fin),
                          "fila": str(numero)})
    wb.close()
    return resultado


def _celda_pintada(celda) -> bool:
    """True si la celda tiene un relleno explícito (no cuenta el estilo vacío)."""
    relleno = celda.fill
    if not relleno or relleno.patternType != "solid":
        return False
    color = relleno.fgColor
    if color.type == "rgb":
        return color.rgb not in (None, "00000000", "FFFFFFFF")
    if color.type == "indexed":
        return color.indexed not in (None, 0, 64, 65)
    if color.type == "theme":
        return color.theme not in (None, 0) or color.tint != 0
    return False


# Convierte una fecha de Excel o texto y registra errores con fila y campo.
def _fecha_temporal(valor, epoca, path, fila, campo):
    from openpyxl.utils.datetime import from_excel
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return ""
    if isinstance(valor, datetime):
        fecha = valor.date()
    elif isinstance(valor, date):
        fecha = valor
    elif isinstance(valor, (int, float)):
        fecha_excel = from_excel(valor, epoca)
        fecha = fecha_excel.date() if isinstance(fecha_excel, datetime) else fecha_excel
    else:
        texto = str(valor).strip()
        fecha = None
        for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
            try:
                fecha = datetime.strptime(texto, formato).date()
                break
            except ValueError:
                continue
        if fecha is None:
            raise ErrorCarga(f"{path}, fila {fila}: fecha inválida en el campo {campo}: {valor!r}.")
    return fecha.isoformat()
