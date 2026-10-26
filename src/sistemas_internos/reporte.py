"""Creación del libro Excel con resumen, observaciones y detalle."""
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# Genera el libro de sistemas internos con resumen, observaciones, detalle y errores.
def generar_reporte(resultado: dict, archivos: dict, salida: str | Path,
                    sistemas: tuple[str, ...] = ("ArgosNet",),
                    campos_por_sistema: dict[str, dict[str, bool]] | None = None) -> Path:
    # Construye un libro con resumen, hojas por sistema y hojas de errores/detalle.
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Revisión de accesos"])
    resumen.append(["Fecha y hora de ejecución", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")])
    resumen.append(["Matriz", archivos["matriz"]])
    fuente_personas = resultado.get("fuente_personas", "Dominio")
    resumen.append([fuente_personas, archivos.get("fuente_personas", archivos["dominio"])])
    resumen.append([f"Registros de {fuente_personas}", len(resultado["dominios"])])
    resumen.append(["Casos no evaluables", sum(d["resultado"].startswith("No evaluable")
                                                for d in resultado["detalle"])])
    resumen.append(["Errores o inconsistencias de datos", len(resultado["errores"])])
    resumen.append([])
    resumen.append(["Registros por sistema", "Archivo TXT", "Registros leídos", "Reglas de matriz",
                    "Logins coincidentes", "Logins sin coincidencia"])
    for sistema in resultado["sistemas"]:
        resumen.append([sistema["nombre"], sistema["archivo"], sistema["registros"],
                        sistema["reglas"], sistema["coincidentes"], sistema["sin_coincidencia"]])
    resumen.append([])
    resumen.append(["Observaciones por categoría", "Cantidad"])
    for categoria, total in sorted(Counter(o["categoria"] for o in resultado["observaciones"]).items()):
        resumen.append([categoria, total])
    resumen.append([])
    if (fuente_personas.strip().casefold() == "salar"
            and archivos.get("dominio_externo_consultado")):
        resumen.append(["Las cuentas sin persona en SALAR y no clasificadas como externas tienen accesos pendientes de revisar o retirar."])
        resumen.append(["Las personas SALAR sin correo y sin coincidencia por nombre en Dominio se identifican como pendientes de creación o revisión en Dominio."])
    else:
        resumen.append([f"Los usuarios fuera de {fuente_personas} tienen accesos pendientes de revisar o retirar."])
    resumen.append(["Las demás observaciones requieren revisión y no confirman por sí solas un acceso indebido."])

    sistemas_reporte = list(dict.fromkeys((*sistemas,
        *(o["sistema"] for o in resultado["observaciones"]))))
    # Separa las observaciones temporales para mostrar sus columnas de vigencia y designación.
    for sistema in sistemas_reporte:
        campos = (campos_por_sistema or {}).get(sistema, {"oficina": True, "gerencia": True})
        incluir_oficina = campos.get("oficina", True)
        incluir_gerencia = campos.get("gerencia", True)
        obs = wb.create_sheet(f"Observaciones {sistema}")
        filas_sistema = [o for o in resultado["observaciones"] if o["sistema"] == sistema]
        temporal_cats = {"Con Acceso Temporal", "Fin Acceso Temporal"}
        encabezados_normales = [
            "Categoría", ("Login usado para el cruce" if fuente_personas.strip().casefold() == "salar"
                          else f"Login ({fuente_personas})"), f"Nombre ({fuente_personas})",
            "Cargo (sistema)", f"Cargo ({fuente_personas})",
        ]
        if fuente_personas.strip().casefold() == "salar":
            encabezados_normales.append("Login SALAR (original)")
        if incluir_oficina:
            encabezados_normales.extend(["Oficina (sistema)", f"Oficina ({fuente_personas})"])
        if incluir_gerencia:
            encabezados_normales.extend(["Gerencia (sistema)", f"Gerencia ({fuente_personas})"])
        encabezados_normales.extend(["Valor esperado (matriz)", "Valor encontrado (sistema)",
                                     "Cargo en matriz"])
        encabezados_temporales = [
            "Categoría", f"Login ({fuente_personas})", f"Nombre ({fuente_personas})",
            "Cargo (sistema)", "Cargo (temporal, actual)",
            "Cargo (temporal, designado)",
        ]
        if incluir_oficina:
            encabezados_temporales.extend(["Oficina (sistema)", "Oficina (temporal, designada)"])
        encabezados_temporales.extend([
            "Valor esperado (matriz)", "Valor encontrado (sistema)",
            "Fecha inicio (temporal)", "Fecha fin (temporal)",
        ])
        # Mantener una tabla por tipo de caso evita columnas temporales vacías
        # en observaciones ordinarias y campos SALAR irrelevantes en temporales.
        filas_normales = []
        for observacion in filas_sistema:
            if observacion["categoria"] in temporal_cats:
                continue
            fila = [observacion["categoria"], observacion["login"], observacion["nombre"],
                observacion["cargo_sistema"], observacion["cargo"],
            ]
            if fuente_personas.strip().casefold() == "salar":
                fila.append(observacion.get("login_salar", ""))
            if incluir_oficina:
                fila.extend([observacion["oficina_sistema"], observacion["oficina_salar"]])
            if incluir_gerencia:
                fila.extend([observacion["gerencia_sistema"], observacion["gerencia_salar"]])
            fila.extend([observacion["esperado"], observacion["encontrado"],
                         observacion.get("cargo_en_matriz")])
            filas_normales.append(fila)
        _tabla(obs, encabezados_normales, filas_normales)
        obs.column_dimensions[get_column_letter(len(encabezados_normales))].hidden = True
        temporales = [o for o in filas_sistema if o["categoria"] in temporal_cats]
        if temporales:
            obs_temporales = wb.create_sheet(f"Temporal {sistema}")
            _tabla(obs_temporales, encabezados_temporales,
                [_fila_temporal(o, incluir_oficina)
                 for o in temporales])
            _colorear_categorias(obs_temporales, temporales, fuente_personas)
        normales = [o for o in filas_sistema if o["categoria"] not in temporal_cats]
        _colorear_categorias(obs, normales, fuente_personas)

    # El detalle y los errores conservan también los casos conformes y no evaluables.
    det = wb.create_sheet("Detalle de comparación")
    _tabla(det, ["Sistema", "Login", "Nombre", "Cargo", "Acceso esperado en Matriz",
                 "Acceso encontrado", "Resultado de la comparación"],
           [[d[k] for k in ("sistema", "login", "nombre", "cargo", "esperados", "encontrados",
                            "resultado")]
            for d in resultado["detalle"]])

    err = wb.create_sheet("Errores de carga")
    _tabla(err, ["Archivo", "Número de fila", "Campo afectado", "Descripción"],
           [[e["archivo"], e["fila"], e["campo"], e["descripcion"]] for e in resultado["errores"]])
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    return salida


# Adapta una observacion temporal a las columnas del reporte.
def _fila_temporal(observacion, incluir_oficina):
    # Adapta una observación temporal al orden de columnas de su hoja de reporte.
    fila = [observacion["categoria"], observacion["login"], observacion["nombre"],
            observacion["cargo_sistema"], observacion["cargo_actual"],
            observacion["cargo_designado"]]
    if incluir_oficina:
        fila.extend([observacion["oficina_sistema"], observacion["oficina_designada"]])
    esperado = observacion["esperado"]
    segmentos = observacion.get("esperado_temporal_segmentos")
    if segmentos:
        esperado = CellRichText([
            TextBlock(InlineFont(color=color), texto) if color else texto
            for texto, color in segmentos
        ])
    fila.extend([esperado, observacion["encontrado"],
                 observacion["fecha_inicio"], observacion["fecha_fin"]])
    return fila


# Escribe encabezados y filas en una hoja de Excel.
def _tabla(ws, encabezados, filas):
    # Escribe encabezados y registros y aplica el formato tabular común.
    ws.append(encabezados)
    for fila in filas:
        ws.append(fila)
    _ajustar_tabla(ws)


# Ajusta anchos, filtros, encabezados y alineacion para facilitar la lectura.
def _ajustar_tabla(ws):
    # Facilita filtrar y leer las hojas: congela encabezados, ajusta anchos y activa ajuste de texto.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="244062")
    for col in ws.columns:
        ancho = min(60, max(12, max((len(str(c.value or "")) for c in col), default=12) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = ancho
    for row in ws.iter_rows(min_row=2):
        for celda in row:
            alineacion = copy(celda.alignment)
            alineacion.wrap_text = True
            alineacion.vertical = "top"
            celda.alignment = alineacion


# Colorea categorias de observacion para distinguir el tipo de caso.
def _colorear_categorias(ws, observaciones, fuente_personas):
    # Aplica un color a cada categoría para distinguir visualmente los tipos de hallazgo.
    colores = {
        "Con Acceso Temporal": "A9D18E",
        "Fin Acceso Temporal": "F4CCCC",
        "Usuario Externo (Dominio)": "D9EAD3",
        "Usuario en Dominio sin registro en sistema": "A9D18E",
        "Usuario sin Dominio pendiente de creación": "FCE4D6",
        (f"Usuarios en {fuente_personas} sin registro en sistema"
         if fuente_personas.strip().casefold() == "salar"
         else f"Usuario en {fuente_personas} sin registro en sistema"): "A9D18E",
        f"Usuario fuera de {fuente_personas} con acceso pendiente de retirar": "F4CCCC",
        "Diferencia de roles": "DDEBF7",
        "Diferencia de Oficina/Gerencia": "FCE4D6",
        "Diferencia de cargo": "FCE4D6",
        "Estado distinto de Activo": "F4CCCC",
        "Registro en sistema sin acceso esperado": "F4CCCC",
        "Cargo sin correspondencia en matriz": "E4DFEC",
        "Duplicados o casos especiales": "D9E1F2",
    }
    for fila, observacion in enumerate(observaciones, 2):
        color = colores.get(observacion["categoria"])
        if color:
            ws.cell(fila, 1).fill = PatternFill("solid", fgColor=color)
