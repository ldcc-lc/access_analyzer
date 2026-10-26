"""Une el análisis de sistemas internos y externos en un único Excel."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analisis_salar_dominio.comparacion_salar_dominio import (
    agregar_hoja_comparacion, comparar_salar_dominio,
)
from src.sistemas_internos.proceso import ejecutar_revision
from src.sistemas_internos.lectores import (DOMINIO_CONFIG, SALAR_CONFIG,
    TemporalesConfig, ErrorCarga, MatrizConfig, PersonasConfig, SistemaConfig)
from .asfi import ASFIConfig, analizar_asfi
from .comun import _leer_matriz_sistema_externo
from .infocred import INFOCREDConfig, analizar_infocred


def ejecutar_revision_completa(matriz: str | Path, personas: str | Path,
        archivos_sistema: dict[str, str | Path], salida: str | Path,
        sistemas: tuple[SistemaConfig, ...], archivo_temporales: str | Path | None = None,
        fuente_personas: str = "SALAR", config_matriz: MatrizConfig = MatrizConfig(),
        config_personas: PersonasConfig = SALAR_CONFIG,
        config_temporales: TemporalesConfig = TemporalesConfig(),
        archivo_dominio: str | Path | None = None,
        config_dominio: PersonasConfig = DOMINIO_CONFIG,
        archivo_asfi: str | Path | None = None,
        config_asfi: ASFIConfig = ASFIConfig(),
        archivo_salar_anterior: str | Path | None = None,
        archivo_infocred: str | Path | None = None,
        config_infocred: INFOCREDConfig = INFOCREDConfig()) -> tuple[Path, dict]:
    """Genera el análisis de los sistemas seleccionados y sus conciliaciones."""
    configuraciones_externas = tuple(config for archivo, config in (
        (archivo_asfi, config_asfi), (archivo_infocred, config_infocred)) if archivo)
    if not sistemas and not configuraciones_externas:
        raise ValueError("Selecciona al menos un sistema interno, ASFI o INFOCRED.")
    # Crea el libro base; también admite una selección compuesta solo por externos.
    generado, resultado_interno = ejecutar_revision(
        matriz, personas, archivos_sistema, salida, sistemas, archivo_temporales,
        fuente_personas, config_matriz, config_personas, config_temporales,
        archivo_dominio, config_dominio,
        incluir_cargos_matriz=bool(
            archivo_dominio and fuente_personas.strip().casefold() == "salar"),
        permitir_sin_sistemas=bool(configuraciones_externas))

    # Si se cargaron ambas fuentes, agrega la conciliación de personas SALAR-Dominio.
    if (archivo_dominio and fuente_personas.strip().casefold() == "salar"):
        cargos_matriz = resultado_interno.pop("_cargos_matriz", set())
        cargos_matriz.update(_cargos_matriz_externos(matriz, configuraciones_externas))
        comparacion_personas = comparar_salar_dominio(
            personas, archivo_dominio, config_personas or SALAR_CONFIG, config_dominio,
            archivo_temporales, config_temporales, archivo_salar_anterior,
            cargos_matriz)
        agregar_hoja_comparacion(generado, comparacion_personas)
        resultado_interno["comparacion_salar_dominio"] = comparacion_personas["observaciones"]
        resultado_interno["detalle_salar_dominio"] = comparacion_personas["detalle"]
        resultado_interno["errores_salar_dominio"] = comparacion_personas["errores"]
        if "bajas_historicas" in comparacion_personas:
            resultado_interno["bajas_historicas"] = comparacion_personas["bajas_historicas"]

    # Los sistemas externos son opcionales; si no hay ninguno, termina aquí.
    if not archivo_asfi and not archivo_infocred:
        resultado_interno["observaciones_totales"] = (
            len(resultado_interno["observaciones"])
            + len(resultado_interno.get("comparacion_salar_dominio", []))
            + len(resultado_interno.get("bajas_historicas", [])))
        return generado, resultado_interno

    errores_previos = {_clave_error(item) for item in resultado_interno["errores"]}
    resultados_externos = []
    if archivo_asfi:
        # Analiza ASFI y conserva el esquema de hojas que ya consume el informe.
        resultado_asfi = analizar_asfi(matriz, personas, archivo_asfi,
            dominio=archivo_dominio, archivo_temporales=archivo_temporales,
            config_temporales=config_temporales, config_dominio=config_dominio,
            config_asfi=config_asfi)
        resultado_asfi["errores"] = [item for item in resultado_asfi["errores"]
            if _clave_error(item) not in errores_previos]
        errores_previos.update(_clave_error(item) for item in resultado_asfi["errores"])
        _agregar_asfi_al_reporte(generado, resultado_asfi, archivo_asfi)
        resultados_externos.append(resultado_asfi)

    if archivo_infocred:
        # INFOCRED comparte la comparación, pero usa sus propios campos y equivalencias.
        resultado_infocred = analizar_infocred(matriz, personas, archivo_infocred,
            dominio=archivo_dominio, archivo_temporales=archivo_temporales,
            config_temporales=config_temporales, config_dominio=config_dominio,
            config_infocred=config_infocred)
        resultado_infocred["errores"] = [item for item in resultado_infocred["errores"]
            if _clave_error(item) not in errores_previos]
        errores_previos.update(_clave_error(item) for item in resultado_infocred["errores"])
        _agregar_sistema_externo_al_reporte(
            generado, resultado_infocred, archivo_infocred, "INFOCRED")
        resultados_externos.append(resultado_infocred)

    combinado = dict(resultado_interno)
    combinado["observaciones"] = (list(resultado_interno["observaciones"])
        + [item for result in resultados_externos for item in result["observaciones"]]
        + [item for result in resultados_externos
           for item in result["observaciones_temporales"]])
    combinado["observaciones_totales"] = (
        len(combinado["observaciones"])
        + len(resultado_interno.get("comparacion_salar_dominio", []))
        + len(resultado_interno.get("bajas_historicas", [])))
    combinado["errores"] = (list(resultado_interno["errores"])
        + [item for result in resultados_externos for item in result["errores"]])
    for result in resultados_externos:
        combinado[result["sistema"].casefold()] = result
    return generado, combinado


def _cargos_matriz_externos(matriz: str | Path,
                            configuraciones: tuple[ASFIConfig | INFOCREDConfig, ...]) -> set[str]:
    """Incluye cargos de los externos seleccionados en la conciliación SALAR-Dominio."""
    if not configuraciones:
        return set()
    try:
        workbook = load_workbook(matriz, read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorCarga(f"No se pudo leer la matriz {matriz}: {exc}") from exc
    try:
        cargos = set()
        for config in configuraciones:
            reglas, _, _ = _leer_matriz_sistema_externo(workbook, config)
            cargos.update(reglas)
        return cargos
    finally:
        workbook.close()


# Agrega hojas, detalle y resumen de un sistema externo al libro de análisis.
def _agregar_sistema_externo_al_reporte(
        path: Path, resultado: dict, archivo_sistema: str | Path,
        nombre_sistema: str) -> None:
    es_asfi = nombre_sistema.casefold() == "asfi"
    observaciones_sheet = f"Observaciones {nombre_sistema}"
    temporal_sheet = f"Temporal {nombre_sistema}"
    detalle_sheet = ("Externos Detalle de comparación" if es_asfi
                     else f"{nombre_sistema} Detalle")
    errores_sheet = ("Externos Errores de carga" if es_asfi
                     else f"{nombre_sistema} Errores")
    try:
        # Conservar el texto enriquecido usado por internos para colorear
        # los accesos esperados en las observaciones temporales.
        workbook = load_workbook(path, rich_text=True)
    except Exception as exc:
        raise ErrorCarga(f"No se pudo actualizar el reporte con {nombre_sistema}: {exc}") from exc
    try:
        _crear_hoja(workbook, observaciones_sheet, [
            "Categoría", "Correo", "Nombre", "Cargo SALAR/Dominio",
            "Valor esperado (matriz)", f"Valor encontrado ({nombre_sistema})",
            "Grupo informe", "Cargo en matriz",
        ], resultado["observaciones"], (
            "categoria", "correo", "nombre", "cargo", "esperado", "encontrado",
            "grupo_informe", "cargo_en_matriz"))
        _crear_hoja(workbook, temporal_sheet, [
            "Categoría", "Correo", "Nombre", "Cargo SALAR", "Cargo de origen",
            "Cargo designado", "Valor esperado (matriz)", f"Valor encontrado ({nombre_sistema})",
            "Fecha inicio", "Fecha fin", "Resultado", "Grupo informe",
            "Cargo en matriz",
        ], resultado["observaciones_temporales"], (
            "categoria", "correo", "nombre", "cargo", "cargo_actual",
            "cargo_designado", "esperado", "encontrado", "fecha_inicio",
            "fecha_fin", "resultado", "grupo_informe", "cargo_en_matriz"))

        for sheet_name in (observaciones_sheet, temporal_sheet):
            sheet = workbook[sheet_name]
            for column in range(sheet.max_column - 1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(column)].hidden = True

        detalle_externo = [{
            "sistema": nombre_sistema, "correo": item.get("correo", ""),
            "nombre": item.get("nombre", ""), "cargo": item.get("cargo", ""),
            "esperado": item.get("esperado", ""),
            "encontrado": item.get("encontrado", ""),
            "resultado": item.get("resultado", ""),
        } for item in resultado["detalle"]]
        _crear_hoja(workbook, detalle_sheet, [
            "Sistema", "Correo", "Nombre", "Cargo", "Acceso esperado en Matriz",
            "Acceso encontrado", "Resultado de la comparación",
        ], detalle_externo, (
            "sistema", "correo", "nombre", "cargo", "esperado", "encontrado", "resultado"))
        errores_externos = []
        for error in resultado["errores"]:
            item = dict(error)
            if error.get("usuario_salar") or error.get("nombre_salar"):
                identidad = (f"Usuario SALAR: {error.get('usuario_salar') or '(vacío)'}; "
                             f"Nombre: {error.get('nombre_salar') or '(vacío)'}. ")
                item["descripcion"] = identidad + str(error.get("descripcion", ""))
            errores_externos.append(item)
        _crear_hoja(workbook, errores_sheet, [
            "Archivo", "Número de fila", "Campo afectado", "Descripción",
        ], errores_externos, ("archivo", "fila", "campo", "descripcion"))

        resumen = workbook["Resumen"]
        errores_row = next((row for row in range(1, resumen.max_row + 1)
            if resumen.cell(row, 1).value == "Errores o inconsistencias de datos"), None)
        if errores_row:
            resumen.cell(errores_row, 2).value = (
                (resumen.cell(errores_row, 2).value or 0) + len(resultado["errores"]))
        resumen.append([])
        resumen.append(["Sistemas externos", nombre_sistema])
        resumen.append([f"Archivo {nombre_sistema}", Path(archivo_sistema).name])
        resumen.append([f"Usuarios {nombre_sistema} consolidados", resultado["usuarios"]])
        resumen.append([f"Usuarios {nombre_sistema} comparados", resultado["usuarios_con_detalle"]])
        resumen.append([f"Observaciones {nombre_sistema}", len(resultado["observaciones"])])
        resumen.append([f"Casos temporales {nombre_sistema}", len(resultado["observaciones_temporales"])])
        resumen.append([f"Errores {nombre_sistema}", len(resultado["errores"])])
        resumen.append([f"Categorías de observación {nombre_sistema}", "Cantidad"])
        for category, count in sorted(Counter(
                item["categoria"] for item in resultado["observaciones"]).items()):
            resumen.append([category, count])
        workbook.save(path)
    finally:
        workbook.close()


# Conserva la función privada original usada para actualizar el reporte de ASFI.
def _agregar_asfi_al_reporte(path: Path, resultado: dict,
                             archivo_asfi: str | Path) -> None:
    _agregar_sistema_externo_al_reporte(path, resultado, archivo_asfi, "ASFI")


# Crea una hoja tabular, escribe registros y aplica formato visual por categoria.
def _crear_hoja(workbook, title, headers, records, fields) -> None:
    # Reemplaza o crea una hoja de datos con formato de tabla y color por categoría.
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for item in records:
        values = []
        for key in fields:
            value = item.get(key, "")
            if key in {"esperado", "encontrado"}:
                value = _ordenar_roles(value)
            values.append(_mostrar(value))
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    category_colors = {
        "Con Acceso Temporal": "A9D18E",
        "Fin Acceso Temporal": "F4CCCC",
        "Diferencia de roles": "DDEBF7",
        "Usuario Externo (Dominio)": "D9EAD3",
        "Usuario sin Dominio pendiente de creación": "FCE4D6",
        "Usuario fuera de SALAR con acceso pendiente de retirar": "F4CCCC",
        "Usuarios en SALAR sin registro en sistema": "A9D18E",
        "Cargo sin correspondencia en matriz": "E4DFEC",
        "Rol - código no coincidentes": "D9E1F2",
        "Duplicados o casos especiales": "D9E1F2",
    }
    for row_index, item in enumerate(records, start=2):
        color = category_colors.get(item.get("categoria"))
        if color:
            sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor=color)
    for column in sheet.columns:
        width = min(60, max(12, max((len(str(cell.value or ""))
                                     for cell in column), default=12) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


# Convierte valores nulos y colecciones en texto seguro para celdas de Excel.
def _mostrar(value) -> str:
    # Convierte valores opcionales y colecciones a texto apto para Excel.
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    return str(value)


# Ordena codigos de acceso para que el reporte sea consistente y facil de comparar.
def _ordenar_roles(value):
    # Ordena roles antes de incorporarlos al reporte combinado.
    if value is None or value == "":
        return value
    if isinstance(value, str):
        roles = [role.strip() for role in value.split(",") if role.strip()]
    else:
        roles = [str(role).strip() for role in value if str(role).strip()]
    return sorted(roles, key=str.casefold)


# Crea la identidad de un error para quitar duplicados entre procesos.
def _clave_error(item: dict) -> tuple:
    # Produce una clave para evitar repetir el mismo error entre análisis.
    return tuple(str(item.get(key, "")) for key in ("archivo", "fila", "campo", "descripcion"))
