"""Genera observaciones de telefonía comparando Dominio con Internos."""
from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ErrorTelefonia(ValueError):
    """Error de lectura o estructura de los archivos de telefonía."""


# Representa una fila de una fuente de telefonia con nombre, telefono y fila de origen.
@dataclass(frozen=True)
class Registro:
    nombre: str
    telefonos: tuple[str, ...]
    fila: int


# Indica las columnas de nombre y telefono que se deben leer.
@dataclass(frozen=True)
class TelefoniaConfig:
    campo_nombre: int
    campo_telefono: int


# Agrupa filas de una persona y consolida sus variantes de nombre y telefonos.
@dataclass
class Grupo:
    clave: str
    nombres: list[str] = field(default_factory=list)
    telefonos: set[str] = field(default_factory=set)
    filas: list[int] = field(default_factory=list)
    registros: int = 0

    # Devuelve el nombre preferido del grupo para presentarlo en el reporte.
    @property
    def nombre(self) -> str:
        return self.nombres[0] if self.nombres else ""


_INDICADORES_GENERICOS = (
    "agencia", "oficina", "gerencia", "departamento", "unidad", "area ",
    "seguridad", "audext", "auditext", "auditor", "call center", "cpd",
    "deposito", "finanzas", "tesoreria", "jefe agencia", "oficial de credito",
    "operaciones", "operador de monitoreo", "monitoreo", "sala central",
    "sala reuniones", "sala de reuniones", "plataforma",
    "centro de", "servicio", "auditoria externa", "agente ia", "comision",
    "banco mundial", "asfi", "supervision", "talento humano", "linea de vida",
)


def _reparar_mojibake(texto: str) -> str:
    """Repara texto UTF-8 interpretado como Windows-1252 cuando es posible."""
    actual = texto
    for _ in range(2):
        if not any(marca in actual for marca in ("Ã", "Â", "â")):
            break
        try:
            candidato = actual.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        marcas = ("Ã", "Â", "â")
        if sum(candidato.count(m) for m in marcas) >= sum(
                actual.count(m) for m in marcas):
            break
        actual = candidato
    return actual.replace("\ufffd", "")


# Normaliza nombres para comparar ignorando acentos, signos y variaciones de texto.
@lru_cache(maxsize=8192)
def normalizar_nombre(nombre: str) -> str:
    nombre = _reparar_mojibake(nombre).strip().casefold()
    nombre = unicodedata.normalize("NFKD", nombre)
    nombre = "".join(c for c in nombre if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", nombre).strip()


# Lee registros de telefonia, conservando la fila y validando los campos configurados.
def _leer_registros(path: str | Path, campos_esperados: int,
                    columna_nombre: int, columna_telefono: int) -> list[Registro]:
    # Lee TXT de telefonía, conserva fila de origen y extrae todos los números del campo.
    archivo = Path(path)
    try:
        raw = archivo.read_bytes()
    except OSError as exc:
        raise ErrorTelefonia(f"No se pudo abrir {archivo}: {exc}") from exc

    texto = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise ErrorTelefonia(f"No se pudo reconocer la codificación de {archivo.name}.")

    registros = []
    for numero, fila in enumerate(csv.reader(texto.splitlines(), delimiter="|"), 1):
        if not fila or (len(fila) == 1 and not fila[0].strip()):
            continue
        if len(fila) < campos_esperados:
            raise ErrorTelefonia(
                f"{archivo.name}, fila {numero}: se esperaban al menos "
                f"{campos_esperados} campos y llegaron {len(fila)}.")
        nombre = _reparar_mojibake(fila[columna_nombre - 1]).strip()
        if not nombre or normalizar_nombre(nombre) in {"null", "none"}:
            continue
        telefono_texto = _reparar_mojibake(fila[columna_telefono - 1]).strip()
        telefonos = tuple(dict.fromkeys(re.findall(r"\d+", telefono_texto)))
        registros.append(Registro(nombre, telefonos, numero))
    if not registros:
        raise ErrorTelefonia(f"{archivo.name}: no se encontraron registros con nombre.")
    return registros


# Agrupa filas con nombres equivalentes y une sus telefonos.
def _agrupar(registros: list[Registro]) -> dict[str, Grupo]:
    # Une filas con el mismo nombre normalizado y consolida sus teléfonos.
    grupos: dict[str, Grupo] = {}
    for registro in registros:
        clave = normalizar_nombre(registro.nombre)
        if not clave:
            continue
        grupo = grupos.setdefault(clave, Grupo(clave))
        if registro.nombre not in grupo.nombres:
            grupo.nombres.append(registro.nombre)
        grupo.telefonos.update(registro.telefonos)
        grupo.filas.append(registro.fila)
        grupo.registros += 1
    return grupos


# Separa un nombre en palabras normalizadas para buscar candidatos.
def _tokens(nombre: str) -> set[str]:
    return set(normalizar_nombre(nombre).split())


def _nombres_equivalentes(nombre_a: str, nombre_b: str) -> bool:
    """Reconoce igualdad, abreviaciones con al menos dos tokens y typos leves."""
    clave_a, clave_b = normalizar_nombre(nombre_a), normalizar_nombre(nombre_b)
    if not clave_a or not clave_b:
        return False
    if clave_a == clave_b:
        return True
    tokens_a, tokens_b = set(clave_a.split()), set(clave_b.split())
    if len(tokens_a) >= 2 and len(tokens_b) >= 2:
        if tokens_a <= tokens_b or tokens_b <= tokens_a:
            return True
        comunes = tokens_a & tokens_b
        if len(comunes) >= 2:
            return SequenceMatcher(None, clave_a, clave_b).ratio() >= 0.9
        # El cálculo difuso es costoso; solo se intenta si al menos una palabra
        # larga coincide y las longitudes son compatibles con un typo menor.
        if (len(comunes) == 1 and len(next(iter(comunes))) >= 4
                and min(len(clave_a), len(clave_b))
                / max(len(clave_a), len(clave_b)) >= 0.82):
            return SequenceMatcher(None, clave_a, clave_b).ratio() >= 0.94
        return False
    if len(tokens_a) == len(tokens_b) == 1:
        return (min(len(clave_a), len(clave_b)) / max(len(clave_a), len(clave_b))
                >= 0.82
                and SequenceMatcher(None, clave_a, clave_b).ratio() >= 0.96)
    return False


# Reconoce etiquetas de area o titulares genericos que no identifican a una persona.
def _es_generico(nombre: str) -> bool:
    clave = normalizar_nombre(nombre)
    if not clave:
        return False
    # Comparar frases y palabras completas evita falsos positivos por
    # subcadenas dentro de apellidos como "Salazar" o "Salas".
    tokens = clave.split()
    frases = [indicador.strip().split() for indicador in _INDICADORES_GENERICOS]
    for frase in frases:
        ancho = len(frase)
        if any(tokens[indice:indice + ancho] == frase
               for indice in range(len(tokens) - ancho + 1)):
            return True
    # AUDEXTERNA, AUDITORES y AUDITORIA EXTERNA son variantes de la misma
    # etiqueta institucional; el prefijo solo se aplica a esos términos.
    return any(token.startswith(("audext", "auditext", "auditor"))
               for token in tokens)


def _grupos_persona_internos(grupos: dict[str, Grupo]) -> list[Grupo]:
    """Une abreviaciones con un solo nombre completo posible en Internos."""
    # Excluye nombres genéricos y une abreviaciones solo si tienen un único nombre completo candidato.
    personas = [grupo for grupo in grupos.values()
                if not any(_es_generico(nombre) for nombre in grupo.nombres)]
    padre = list(range(len(personas)))

    # Encuentra la raiz de un conjunto en la estructura usada para unir abreviaciones.
    def raiz(indice: int) -> int:
        while padre[indice] != indice:
            padre[indice] = padre[padre[indice]]
            indice = padre[indice]
        return indice

    tokens = [_tokens(grupo.nombre) for grupo in personas]
    candidatos = [[] for _ in personas]
    for i, tokens_i in enumerate(tokens):
        if len(tokens_i) < 2:
            continue
        for j, tokens_j in enumerate(tokens):
            if i == j or len(tokens_j) <= len(tokens_i):
                continue
            if tokens_i <= tokens_j:
                candidatos[i].append(j)
    for a, indices in enumerate(candidatos):
        if len(indices) == 1:
            b = indices[0]
            ra, rb = raiz(a), raiz(b)
            if ra != rb:
                padre[rb] = ra

    conjuntos: dict[int, Grupo] = {}
    for indice, grupo in enumerate(personas):
        destino = conjuntos.setdefault(raiz(indice), Grupo(grupo.clave))
        for nombre in grupo.nombres:
            if nombre not in destino.nombres:
                destino.nombres.append(nombre)
        destino.telefonos.update(grupo.telefonos)
        destino.filas.extend(grupo.filas)
        destino.registros += grupo.registros
    return list(conjuntos.values())


# Indexa cada telefono con las personas o grupos que lo tienen asignado.
def _mapa_telefonos(grupos: dict[str, Grupo]) -> dict[str, list[Grupo]]:
    # Indexa titulares por teléfono para detectar números compartidos o asignados a otra persona.
    mapa: dict[str, list[Grupo]] = defaultdict(list)
    for grupo in grupos.values():
        for telefono in grupo.telefonos:
            mapa[telefono].append(grupo)
    return mapa


# Devuelve nombres asociados a un conjunto de telefonos sin repetirlos.
def _nombres_por_telefonos(telefonos: set[str], mapa: dict[str, list[Grupo]]) -> list[str]:
    nombres = []
    vistos = set()
    for telefono in sorted(telefonos):
        for grupo in mapa.get(telefono, []):
            for nombre in grupo.nombres:
                clave = normalizar_nombre(nombre)
                if clave not in vistos:
                    vistos.add(clave)
                    nombres.append(nombre)
    return nombres


# Indexa grupos por palabras de sus nombres para acelerar las busquedas.
def _indice_nombres(grupos: dict[str, Grupo]) -> dict[str, set[str]]:
    indice: dict[str, set[str]] = defaultdict(set)
    for grupo in grupos.values():
        for nombre in grupo.nombres:
            for token in _tokens(nombre):
                indice[token].add(grupo.clave)
    return indice


# Reduce candidatos por tokens y confirma equivalencia de nombres.
def _buscar_nombres(nombre: str, grupos: dict[str, Grupo],
                    indice: dict[str, set[str]]) -> list[Grupo]:
    # Reduce candidatos por tokens y luego verifica equivalencia de nombres.
    posibles = set()
    for token in _tokens(nombre):
        posibles.update(indice.get(token, ()))
    return [grupos[clave] for clave in posibles
            if any(_nombres_equivalentes(nombre, otro)
                   for otro in grupos[clave].nombres)]


# Construye una observacion de telefonia con sus dos fuentes y numeros.
def _fila(categoria: str, nombre_dominio: str = "", nombre_isabel: str = "",
          telefono_dominio: str = "", telefono_isabel: str = "") -> dict[str, str]:
    return {
        "categoria": categoria,
        "nombre_dominio": nombre_dominio,
        "nombre_isabel": nombre_isabel,
        "telefono_dominio": telefono_dominio,
        "telefono_isabel": telefono_isabel,
    }


def _nombres_exactamente_iguales(nombre_dominio: str, nombre_isabel: str) -> bool:
    """Compara nombres normalizados, sin aceptar abreviaciones ni similitud difusa."""
    dominio = {normalizar_nombre(nombre) for nombre in nombre_dominio.split("|")}
    isabel = {normalizar_nombre(nombre) for nombre in nombre_isabel.split("|")}
    dominio.discard("")
    isabel.discard("")
    return len(dominio) == 1 and dominio == isabel


# Aplica reglas para nombres genericos, numeros compartidos, altas y bajas.
def _analizar(dominio_registros: list[Registro], internos_registros: list[Registro]):
    # Aplica reglas de titulares, números faltantes, duplicados y bajas en orden de prioridad.
    dominio = _agrupar(dominio_registros)
    internos = _agrupar(internos_registros)
    mapa_dominio = _mapa_telefonos(dominio)
    mapa_internos = _mapa_telefonos(internos)
    indice_nombres_dominio = _indice_nombres(dominio)
    indice_nombres_internos = _indice_nombres(internos)
    observaciones = []

    doble_dominio_exactos = 0
    doble_dominio_revision = 0
    # Separa coincidencias exactas por nombre de los candidatos que requieren revision.
    for persona in dominio.values():
        if len(persona.telefonos) < 2 or any(
                _es_generico(nombre) for nombre in persona.nombres):
            continue
        nombres_dominio = {normalizar_nombre(nombre) for nombre in persona.nombres}
        exactos = [grupo for grupo in internos.values()
                   if any(normalizar_nombre(nombre) in nombres_dominio
                          for nombre in grupo.nombres)
                   and any(not _es_generico(nombre) for nombre in grupo.nombres)]
        candidatos = [
            grupo for nombre in persona.nombres
            for grupo in _buscar_nombres(nombre, internos, indice_nombres_internos)
            if any(not _es_generico(candidato) for candidato in grupo.nombres)
        ]
        candidatos = list({grupo.clave: grupo for grupo in candidatos}.values())
        if len(exactos) == 1:
            categoria = "Doble interno Dominio - Coincidencia exacta por nombre"
            grupos_isabel = exactos
            doble_dominio_exactos += 1
        else:
            categoria = "Doble interno Dominio - Revision manual"
            grupos_isabel = candidatos
            doble_dominio_revision += 1
        observaciones.append(_fila(
            categoria,
            " | ".join(persona.nombres),
            " | ".join(dict.fromkeys(
                nombre for grupo in grupos_isabel for nombre in grupo.nombres)),
            ", ".join(sorted(persona.telefonos)),
            ", ".join(sorted({telefono for grupo in grupos_isabel
                               for telefono in grupo.telefonos}))))
    for grupo in internos.values():
        if not any(_es_generico(nombre) for nombre in grupo.nombres):
            continue
        nombres_dom = _nombres_por_telefonos(grupo.telefonos, mapa_dominio)
        observaciones.append(_fila(
            "Nombres Genéricos Isabel", " | ".join(nombres_dom),
            " | ".join(grupo.nombres),
            ", ".join(sorted({t for t in grupo.telefonos if t in mapa_dominio})),
            ", ".join(sorted(grupo.telefonos))))

    # 2. Personas de Internos con más de un número, agrupando abreviaciones
    # inequívocas como "Marcelo Lopez" y "Rene Marcelo Lopez Sanchez".
    personas = _grupos_persona_internos(internos)
    for persona in personas:
        if len(persona.telefonos) < 2:
            continue
        nombres_dom = _nombres_por_telefonos(persona.telefonos, mapa_dominio)
        observaciones.append(_fila(
            "Usuarios con doble interno Isabel", " | ".join(nombres_dom),
            " | ".join(persona.nombres),
            ", ".join(sorted({t for t in persona.telefonos if t in mapa_dominio})),
            ", ".join(sorted(persona.telefonos))))

    # 3. El teléfono es la clave para detectar titulares diferentes.
    todos_telefonos = sorted(set(mapa_dominio) & set(mapa_internos))
    for telefono in todos_telefonos:
        grupos_dom, grupos_int = mapa_dominio[telefono], mapa_internos[telefono]
        nombres_dom = [n for g in grupos_dom for n in g.nombres]
        nombres_int = [n for g in grupos_int for n in g.nombres]
        if not nombres_dom or not nombres_int:
            continue
        if all(_nombres_equivalentes(nd, ni)
               for nd in nombres_dom for ni in nombres_int):
            continue
        observaciones.append(_fila(
            "Teléfono con diferente usuario",
            " | ".join(dict.fromkeys(nombres_dom)),
            " | ".join(dict.fromkeys(nombres_int)), telefono, telefono))

    # 4. Números asignados a personas activas de Dominio que todavía no están
    #    registrados en Isabel. Los números ya existentes en Isabel se dejan
    #    para la categoría de titular diferente, evitando duplicar alertas.
    todos_telefonos_isabel = set(mapa_internos)
    for persona in dominio.values():
        if any(_es_generico(nombre) for nombre in persona.nombres):
            continue
        telefonos_faltantes = persona.telefonos - todos_telefonos_isabel
        if not telefonos_faltantes:
            continue
        personas_isabel = [
            grupo for nombre in persona.nombres
            for grupo in _buscar_nombres(nombre, internos, indice_nombres_internos)
        ]
        nombres_isabel = list(dict.fromkeys(
            nombre for grupo in personas_isabel for nombre in grupo.nombres))
        telefonos_isabel_persona = set(
            telefono for grupo in personas_isabel for telefono in grupo.telefonos)
        observaciones.append(_fila(
            "Dar de Alta en Isabel", " | ".join(persona.nombres),
            " | ".join(nombres_isabel),
            ", ".join(sorted(telefonos_faltantes)),
            ", ".join(sorted(telefonos_isabel_persona))))

    # 5. Si la persona de Internos ya no aparece en Dominio, sugerir su baja.
    # Los genéricos se excluyen porque no son usuarios personales.
    for persona in personas:
        nombres_en_dominio = [
            grupo for nombre in persona.nombres
            for grupo in _buscar_nombres(nombre, dominio, indice_nombres_dominio)
        ]
        if nombres_en_dominio:
            continue
        nombres_dom = _nombres_por_telefonos(persona.telefonos, mapa_dominio)
        observaciones.append(_fila(
            "Dar de Baja en Isabel", " | ".join(nombres_dom),
            " | ".join(persona.nombres),
            ", ".join(sorted({t for t in persona.telefonos if t in mapa_dominio})),
            ", ".join(sorted(persona.telefonos))))

    for observacion in observaciones:
        categoria = observacion["categoria"].casefold()
        if "doble interno dominio - coincidencia exacta" in categoria:
            coincidencia = "Coincidencia exacta por nombre"
        elif "doble interno dominio - revision manual" in categoria:
            coincidencia = "Revision manual"
        elif categoria in {"telefono con diferente usuario", "nombres genericos isabel"}:
            coincidencia = "Revision manual"
        elif _nombres_exactamente_iguales(
                observacion["nombre_dominio"], observacion["nombre_isabel"]):
            coincidencia = "Coincidencia exacta por nombre"
        else:
            coincidencia = "Revision manual"
        observacion["clasificacion_nombre"] = coincidencia

    resumen = {
        "observaciones": len(observaciones),
        "coincidencias_exactas_nombre": sum(
            item["clasificacion_nombre"] == "Coincidencia exacta por nombre"
            for item in observaciones),
        "revision_manual": sum(
            item["clasificacion_nombre"] == "Revision manual"
            for item in observaciones),
        "nombres_duplicados": sum(
            1 for persona in personas if len(persona.telefonos) > 1),
        "usuarios_doble_interno_dominio": (
            doble_dominio_exactos + doble_dominio_revision),
        "usuarios_doble_interno_dominio_exactos": doble_dominio_exactos,
        "usuarios_doble_interno_dominio_revision_manual": doble_dominio_revision,
        "ambiguas": 0,
    }
    return resumen, observaciones


def _inventario_completo(dominio_registros: list[Registro],
                         internos_registros: list[Registro]) -> list[dict[str, str]]:
    """Lista todas las identidades de ambas fuentes y sus posibles cruces."""
    dominio = _agrupar(dominio_registros)
    internos = _agrupar(internos_registros)
    exactas = sorted(set(dominio) & set(internos))
    emparejadas = set(exactas)
    telefonos_isabel = set().union(*(grupo.telefonos for grupo in internos.values()))
    filas = []

    def valores(grupo):
        return {
            "nombre": " | ".join(grupo.nombres),
            "telefonos": ", ".join(sorted(grupo.telefonos)),
            "filas": ", ".join(map(str, sorted(grupo.filas))),
        }

    def registro(estado, accion="", grupo_dominio=None, grupo_isabel=None,
                 candidatos=""):
        d = valores(grupo_dominio) if grupo_dominio else {}
        i = valores(grupo_isabel) if grupo_isabel else {}
        return {
            "estado": estado,
            "accion": accion,
            "nombre_dominio": d.get("nombre", ""),
            "telefonos_dominio": d.get("telefonos", ""),
            "filas_dominio": d.get("filas", ""),
            "nombre_isabel": i.get("nombre", ""),
            "telefonos_isabel": i.get("telefonos", ""),
            "filas_isabel": i.get("filas", ""),
            "candidatos": candidatos,
        }

    for clave in exactas:
        persona = dominio[clave]
        faltantes = persona.telefonos - telefonos_isabel
        accion = "Dar de Alta en Isabel" if faltantes else ""
        if faltantes:
            accion += ": " + ", ".join(sorted(faltantes))
        filas.append(registro("Coincidencia exacta por nombre", accion,
                              persona, internos[clave]))

    dominio_sin_cruce = {key: group for key, group in dominio.items()
                         if key not in emparejadas}
    internos_sin_cruce = {key: group for key, group in internos.items()
                          if key not in emparejadas}
    indice_internos = _indice_nombres(internos)
    indice_dominio = _indice_nombres(dominio)
    persona_interna_por_alias = {
        normalizar_nombre(nombre): persona
        for persona in _grupos_persona_internos(internos)
        for nombre in persona.nombres
    }

    for clave, group in sorted(dominio_sin_cruce.items()):
        candidatos = _buscar_nombres(group.nombre, internos, indice_internos)
        descripcion = " | ".join(
            f"{candidate.nombre} [{', '.join(sorted(candidate.telefonos))}]"
            for candidate in candidatos)
        estado = "Revision manual" if candidatos else "Solo en Dominio"
        faltantes = group.telefonos - telefonos_isabel
        accion = "Dar de Alta en Isabel" if faltantes else ""
        if faltantes:
            accion += ": " + ", ".join(sorted(faltantes))
        filas.append(registro(estado, accion, grupo_dominio=group,
                              candidatos=descripcion))

    for clave, group in sorted(internos_sin_cruce.items()):
        es_generico = any(_es_generico(nombre) for nombre in group.nombres)
        persona = persona_interna_por_alias.get(clave, group)
        candidatos = ([] if es_generico else list({candidate.clave: candidate
            for nombre in persona.nombres
            for candidate in _buscar_nombres(nombre, dominio, indice_dominio)
        }.values()))
        descripcion = " | ".join(
            f"{candidate.nombre} [{', '.join(sorted(candidate.telefonos))}]"
            for candidate in candidatos)
        estado = ("Titular generico en Isabel" if es_generico else
                  "Revision manual" if candidatos else "Solo en Isabel")
        accion = "Dar de Baja en Isabel" if not candidatos and not es_generico else ""
        filas.append(registro(estado, accion, grupo_isabel=group,
                              candidatos=descripcion))

    return filas


def generar_reporte_telefonia(dominio_path: str | Path, internos_path: str | Path,
                              salida: str | Path,
                              config_dominio: TelefoniaConfig = TelefoniaConfig(1, 11),
                              config_isabel: TelefoniaConfig = TelefoniaConfig(1, 2)
                              ) -> tuple[Path, dict]:
    """Crea una hoja de observaciones; Isabel (Internos) es la fuente oficial."""
    # Dominio representa usuarios activos y la lista Internos es el inventario telefónico revisado.
    dominio = _leer_registros(
        dominio_path, max(config_dominio.campo_nombre, config_dominio.campo_telefono),
        config_dominio.campo_nombre, config_dominio.campo_telefono)
    internos = _leer_registros(
        internos_path, max(config_isabel.campo_nombre, config_isabel.campo_telefono),
        config_isabel.campo_nombre, config_isabel.campo_telefono)
    resumen, observaciones = _analizar(dominio, internos)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Observaciones"
    encabezados = [
        "Categoria", "Clasificacion por nombre", "Nombre (Dominio)", "Telefono (Dominio)",
        "Nombre (Isabel)", "Telefono (Isabel)",
    ]
    campos = ("categoria", "clasificacion_nombre", "nombre_dominio",
              "telefono_dominio", "nombre_isabel", "telefono_isabel")
    _escribir_observaciones(hoja, encabezados, observaciones, campos)
    exactas = libro.create_sheet("Coincidencia exacta")
    _escribir_observaciones(exactas, encabezados,
        [item for item in observaciones
         if item["clasificacion_nombre"] == "Coincidencia exacta por nombre"], campos)
    manuales = libro.create_sheet("Revision manual")
    _escribir_observaciones(manuales, encabezados,
        [item for item in observaciones
         if item["clasificacion_nombre"] == "Revision manual"], campos)
    _escribir_inventario(libro.create_sheet("Inventario completo"),
        _inventario_completo(dominio, internos))

    destino = Path(salida)
    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino, resumen


def _escribir_observaciones(hoja, encabezados, observaciones, campos):
    hoja.append(encabezados)
    for observacion in observaciones:
        hoja.append([observacion.get(campo, "") for campo in campos])
    if not observaciones:
        hoja.append(["Sin observaciones"] + [""] * (len(encabezados) - 1))
    _formatear_hoja(hoja)


def _escribir_inventario(hoja, registros):
    encabezados = [
        "Estado del cruce", "Observacion asociada", "Nombre (Dominio)", "Telefonos (Dominio)",
        "Filas de origen (Dominio)", "Nombre (Isabel)", "Telefonos (Isabel)",
        "Filas de origen (Isabel)", "Candidatos para revision manual",
    ]
    hoja.append(encabezados)
    for item in registros:
        hoja.append([item.get(key, "") for key in (
            "estado", "accion", "nombre_dominio", "telefonos_dominio",
            "filas_dominio", "nombre_isabel", "telefonos_isabel",
            "filas_isabel", "candidatos")])
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    anchos = (31, 34, 42, 24, 22, 42, 24, 22, 70)
    for indice, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="244062")
    colores = {
        "Coincidencia exacta por nombre": "D9EAD3",
        "Revision manual": "FCE4D6",
        "Solo en Dominio": "DDEBF7",
        "Solo en Isabel": "F4CCCC",
    }
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        color = colores.get(str(fila[0].value or ""))
        if color:
            fila[0].fill = PatternFill("solid", fgColor=color)


# Aplica formato, filtros, anchos y colores a la hoja de observaciones.
def _formatear_hoja(hoja):
    # Aplica filtros, anchos, encabezados y colores por tipo de observación.
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    anchos = (48, 32, 42, 24, 42, 24)
    for indice, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="244062")
    colores_categoria = {
        "dar de alta en isabel": "D9EAD3",
        "dar de baja en isabel": "F4CCCC",
        "nombres genericos isabel": "E4DFEC",
        "telefono con diferente usuario": "FCE4D6",
        "usuarios con doble interno isabel": "DDEBF7",
        "doble interno dominio coincidencia exacta por nombre": "D9EAD3",
        "doble interno dominio revision manual": "FCE4D6",
    }
    colores_coincidencia = {
        "coincidencia exacta por nombre": "D9EAD3",
        "revision manual": "FCE4D6",
    }
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        categoria = normalizar_nombre(str(fila[0].value or ""))
        color = colores_categoria.get(categoria)
        if color:
            fila[0].fill = PatternFill("solid", fgColor=color)
        clasificacion = normalizar_nombre(str(fila[1].value or ""))
        color_clasificacion = colores_coincidencia.get(clasificacion)
        if color_clasificacion:
            fila[1].fill = PatternFill("solid", fgColor=color_clasificacion)
