"""Extracción de usuarios y roles del reporte PDF de ASFI.

La extracción usa la posición de las palabras dentro de las columnas del PDF.
No necesita una lista predefinida de roles: los nombres nuevos se leen igual
que los que ya conocemos.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


_PATRON_CORREO = re.compile(
    r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE
)
_PATRON_FECHA = re.compile(r"\b\d{4}/\d{2}/\d{2}\b")
_PATRON_PAGINA = re.compile(r"^p[aá]gina\s+\d+\s+de\s+\d+$", re.IGNORECASE)
_SEPARACION_MAX_CONTINUACION_ROL = 14


# Quita saltos y espacios repetidos antes de interpretar palabras del PDF.
def _normalizar_espacios(texto: str) -> str:
    return re.sub(r"\s+", " ", texto or "").strip()


def _agrupar_palabras_por_linea(palabras: list[dict[str, Any]], tolerancia: float = 2.5):
    """Agrupa palabras que están en la misma línea visual del PDF."""
    lineas: list[dict[str, Any]] = []
    for palabra in sorted(palabras, key=lambda item: (item["top"], item["x0"])):
        if not lineas or abs(palabra["top"] - lineas[-1]["top"]) > tolerancia:
            lineas.append({"top": palabra["top"], "palabras": [palabra]})
        else:
            lineas[-1]["palabras"].append(palabra)
    return [linea["palabras"] for linea in lineas]


def _limites_columna_roles(lineas: list[list[dict[str, Any]]]):
    """Calcula los bordes de ROLES a partir de SISTEMA, ROLES y FECHA ALTA."""
    for linea in lineas:
        palabras = {p["text"].strip().upper(): p for p in linea}
        sistema = palabras.get("SISTEMA")
        roles = palabras.get("ROLES")
        fecha = next(
            (p for p in linea if p["text"].strip().upper().startswith("FECHA")),
            None,
        )
        if not (sistema and roles and fecha):
            continue

        centro_sistema = (sistema["x0"] + sistema["x1"]) / 2
        centro_roles = (roles["x0"] + roles["x1"]) / 2
        alta = next(
            (p for p in linea if p["text"].strip().upper() == "ALTA"), None
        )
        centro_fecha = ((fecha["x0"] + alta["x1"]) / 2 if alta else
                        (fecha["x0"] + fecha["x1"]) / 2)
        # Los puntos medios separan columnas aunque los datos estén alineados
        # de forma distinta a los títulos (la fecha de cada rol se alinea a la derecha).
        inicio_roles = (centro_sistema + centro_roles) / 2
        fin_roles = (centro_roles + centro_fecha) / 2
        return inicio_roles, fin_roles
    return None


def _extraer_nombre(texto_linea: str) -> str:
    """Obtiene el nombre ubicado entre las etiquetas Nombre y Doc. Ident."""
    match = re.search(
        r"\bNombre\s*:\s*(.*?)(?=\s*Doc\.?\s*Ident\.?\s*:|$)",
        texto_linea,
        flags=re.IGNORECASE,
    )
    return _normalizar_espacios(match.group(1)) if match else ""


def _extraer_nombre_cercano(
    palabras_pagina: list[dict[str, Any]], referencia_y: float
) -> str:
    """Lee el nombre entre sus etiquetas incluso si quedó en otra línea visual."""
    ventana = [p for p in palabras_pagina if abs(p["top"] - referencia_y) <= 9]
    etiqueta_nombre = next(
        (p for p in ventana if p["text"].strip().casefold().startswith("nombre")), None
    )
    etiqueta_doc = next(
        (p for p in ventana if p["text"].strip().casefold().startswith("doc")
         and etiqueta_nombre and p["x0"] > etiqueta_nombre["x0"]),
        None,
    )
    if not etiqueta_nombre or not etiqueta_doc:
        return ""

    palabras_nombre = [
        p for p in ventana
        if etiqueta_nombre["x1"] <= p["x0"]
        and p["x1"] <= etiqueta_doc["x0"]
    ]
    return _normalizar_espacios(
        " ".join(p["text"] for p in sorted(palabras_nombre,
                                             key=lambda item: (item["top"], item["x0"])))
    )


# Descarta encabezados, pies de pagina y texto que no es un registro.
def _es_linea_de_encabezado_o_pie(texto: str) -> bool:
    normalizado = _normalizar_espacios(texto).casefold()
    return (
        normalizado in {"roles", "sistema", "fecha alta"}
        or all(etiqueta in normalizado.split() for etiqueta in ("sistema", "roles"))
        or bool(_PATRON_PAGINA.fullmatch(normalizado))
        or normalizado.startswith("reporte de usuarios y roles vigentes")
    )


# Crea la estructura inicial para consolidar las filas de un correo.
def _nueva_persona(personas: dict[str, dict[str, Any]], correo: str):
    clave = correo.casefold()
    return personas.setdefault(
        clave, {"Correo": correo, "Nombre": "", "Roles": []}
    )


# Agrega un rol a una persona y une fragmentos de rol continuados en el PDF.
def _agregar_rol(persona: dict[str, Any], rol: str, continuacion: bool = False):
    rol = _normalizar_espacios(rol)
    if not rol:
        return

    if continuacion and persona["Roles"]:
        combinado = _normalizar_espacios(f"{persona['Roles'][-1]} {rol}")
        persona["Roles"][-1] = combinado
        return

    if rol.casefold() not in {existente.casefold() for existente in persona["Roles"]}:
        persona["Roles"].append(rol)


# Recorre paginas y posiciones de palabras para reconstruir usuarios y roles ASFI.
def _extraer_personas(pdf) -> dict[str, dict[str, Any]]:
    personas: dict[str, dict[str, Any]] = {}
    ultima_fila_rol: dict[str, float] = {}
    limites = None
    clave_actual: str | None = None

    for pagina in pdf.pages:
        palabras_pagina = pagina.extract_words()
        lineas = _agrupar_palabras_por_linea(palabras_pagina)
        limites_pagina = _limites_columna_roles(lineas)
        if limites_pagina:
            limites = limites_pagina
        if not limites:
            continue

        inicio_roles, fin_roles = limites
        for palabras in lineas:
            texto_linea = " ".join(p["text"] for p in palabras)
            correo_match = _PATRON_CORREO.search(texto_linea)
            if correo_match:
                correo = correo_match.group(0).strip()
                persona = _nueva_persona(personas, correo)
                clave_actual = correo.casefold()
                nombre = _extraer_nombre_cercano(palabras_pagina, palabras[0]["top"])
                if not nombre:
                    nombre = _extraer_nombre(texto_linea)
                if nombre and not persona["Nombre"]:
                    persona["Nombre"] = nombre
                continue

            if clave_actual is None:
                continue
            persona = personas[clave_actual]

            # Algunos PDFs separan las etiquetas de usuario y nombre en líneas.
            if "Nombre" in texto_linea and not persona["Nombre"]:
                nombre = _extraer_nombre_cercano(palabras_pagina, palabras[0]["top"])
                if not nombre:
                    nombre = _extraer_nombre(texto_linea)
                if nombre:
                    persona["Nombre"] = nombre
                continue

            if _es_linea_de_encabezado_o_pie(texto_linea):
                continue

            palabras_rol = [
                p for p in palabras if inicio_roles <= p["x0"] < fin_roles
            ]
            texto_rol = _normalizar_espacios(" ".join(p["text"] for p in palabras_rol))
            if not texto_rol:
                continue

            # Una fecha identifica una nueva asignación. Una línea de rol sin
            # fecha contigua se trata como continuación del rol anterior.
            tiene_fecha = any(
                _PATRON_FECHA.search(p["text"])
                for p in palabras
                if p["x0"] >= fin_roles
            )
            if tiene_fecha:
                _agregar_rol(persona, texto_rol)
                ultima_fila_rol[clave_actual] = palabras[0]["top"]
            elif (persona["Roles"] and clave_actual in ultima_fila_rol
                  and palabras[0]["top"] - ultima_fila_rol[clave_actual]
                  <= _SEPARACION_MAX_CONTINUACION_ROL):
                _agregar_rol(persona, texto_rol, continuacion=True)

    if not personas:
        raise ValueError(
            "No se encontraron usuarios en el PDF. Verifica que tenga texto seleccionable "
            "y que incluya las columnas SISTEMA, ROLES y FECHA ALTA."
        )
    return personas


def extraer_roles_asfi(ruta_pdf: str | Path) -> list[dict[str, str]]:
    """Extrae una fila por usuario y rol; conserva usuarios sin roles con rol vacío."""
    ruta = Path(ruta_pdf)
    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró el archivo PDF: {ruta}")
    if ruta.suffix.casefold() != ".pdf":
        raise ValueError(f"El archivo debe ser PDF: {ruta}")

    try:
        import pdfplumber
    except ImportError as error:
        raise RuntimeError(
            "Falta pdfplumber. Instala las dependencias del proyecto para leer PDFs."
        ) from error

    with pdfplumber.open(ruta) as pdf:
        personas = _extraer_personas(pdf)

    filas: list[dict[str, str]] = []
    for persona in personas.values():
        for rol in persona["Roles"] or [""]:
            filas.append(
                {"Correo": persona["Correo"], "Nombre": persona["Nombre"], "Rol": rol}
            )
    return filas


def convertir_pdf_a_excel(
    ruta_pdf: str | Path, ruta_excel: str | Path | None = None
) -> Path:
    """Extrae el reporte ASFI y guarda Correo, Nombre y Rol en un Excel."""
    ruta = Path(ruta_pdf)
    filas = extraer_roles_asfi(ruta)
    destino = Path(ruta_excel) if ruta_excel else ruta.with_suffix(".xlsx")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError(
            "Falta openpyxl. Instala las dependencias del proyecto para crear Excel."
        ) from error

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Roles ASFI"
    hoja.append(["Correo", "Nombre", "Rol"])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
    for fila in filas:
        hoja.append([fila["Correo"], fila["Nombre"], fila["Rol"]])

    hoja.freeze_panes = "A2"
    hoja.column_dimensions["A"].width = 36
    hoja.column_dimensions["B"].width = 42
    hoja.column_dimensions["C"].width = 58
    if filas:
        tabla = Table(displayName="RolesASFI", ref=hoja.dimensions)
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        hoja.add_table(tabla)
    else:
        # Sin tabla no existe el autofiltro que esta suele aportar.
        hoja.auto_filter.ref = hoja.dimensions

    _aplicar_formato_roles(hoja)

    libro.save(destino)
    return destino


def _aplicar_formato_roles(hoja) -> None:
    """Mantiene encabezado azul y filtros, con filas blancas y bordes negros."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import TableStyleInfo

    borde = Side(style="thin", color="FF000000")
    cuadricula = Border(left=borde, right=borde, top=borde, bottom=borde)
    for row in hoja.iter_rows():
        for cell in row:
            cell.border = cuadricula
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="FF1F4E78")
                cell.font = Font(bold=True, color="FFFFFF")
            else:
                cell.fill = PatternFill("solid", fgColor="FFFFFFFF")
    for tabla in hoja.tables.values():
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=False,
            showColumnStripes=False)
