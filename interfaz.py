"""Interfaz de escritorio para seleccionar archivos y ejecutar la revisión."""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
from dataclasses import fields, replace
from datetime import date
from pathlib import Path

from PySide6.QtCore import QSettings, QStandardPaths, QThread, Signal, Qt, QTimer
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSplitter,
    QStackedWidget,
    QStyle,
    QVBoxLayout,
    QWidget,
)
from PySide6.QtGui import QFontMetrics, QIcon, QPixmap

from main import SISTEMAS
from src.sistemas_internos.lectores import (
    DOMINIO_CONFIG,
    SALAR_CONFIG,
    MatrizConfig,
    PersonasConfig,
    SistemaConfig,
    TemporalesConfig,
)
from src.sistemas_externos.proceso import ejecutar_revision_completa
from src.sistemas_externos.asfi import ASFIConfig
from src.sistemas_externos.comun import SistemaExternoConfig
from src.sistemas_externos.infocred import INFOCREDConfig
from src.sistema_telefonia import TelefoniaConfig, generar_reporte_telefonia
from src.conversor_pdf_excel import (
    convertir_pdf_a_excel,
    convertir_pdf_sintesis_a_excel,
)
from src.informe.generar_informe import DEFAULT_TEMPLATE, generar_informe


ROOT = Path(__file__).resolve().parent
MATRIX_FILTER = "Archivos Excel (*.xlsx *.xlsm)"
TXT_FILTER = "Archivos de texto (*.txt);;Todos los archivos (*)"
APP_USER_MODEL_ID = "BancoDesarrolloProductivo.AnalizadorDeAccesos"
PDF_EXCEL_NAV_TEXT = "▧ Conversor\n  PDF a Excel"


def recurso_aplicacion(relative_path: str) -> Path:
    """Resuelve recursos tanto desde el proyecto como desde PyInstaller."""
    base = Path(getattr(sys, "_MEIPASS", ROOT))
    return base / relative_path


def _crear_ruta_temporal_excel(prefijo: str) -> Path:
    """Reserva un nombre temporal XLSX que los generadores pueden sobrescribir."""
    descriptor, ruta = tempfile.mkstemp(prefix=prefijo, suffix=".xlsx")
    os.close(descriptor)
    return Path(ruta)


# Widgets reutilizables para selección de archivos, rutas y configuraciones.
class DropPathEdit(QLineEdit):
    """Zona compacta para arrastrar un archivo; nunca muestra su ruta."""

    fileChanged = Signal(str)
    dropRejected = Signal(str)

    def __init__(self, extensions: set[str], parent=None):
        super().__init__(parent)
        self.extensions = {extension.casefold() for extension in extensions}
        self._file_path = ""
        self.setObjectName("filePicker")
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setMinimumWidth(100)
        self.setMinimumHeight(38)
        self.setPlaceholderText("⇧   Arrastra el archivo aquí")
        self.setToolTip("Arrastra un archivo aquí o usa el botón Buscar.")
        self.setReadOnly(True)
        self.setAcceptDrops(True)

    def set_file_path(self, path: str):
        self._file_path = path or ""
        super().setText("✓" if self._file_path else "")
        self.setProperty("fileLoaded", bool(self._file_path))
        self.setToolTip(
            "Archivo cargado. Usa el basurero para quitarlo."
            if self._file_path
            else "Arrastra un archivo aquí o usa el botón Buscar."
        )
        self.style().unpolish(self)
        self.style().polish(self)
        self.update()
        self.fileChanged.emit(self._file_path)

    def file_path(self) -> str:
        return self._file_path

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        paths = [
            Path(url.toLocalFile())
            for url in event.mimeData().urls()
            if url.isLocalFile() and url.toLocalFile()
        ]
        if len(paths) != 1:
            self.dropRejected.emit("Arrastra un solo archivo a la vez.")
            event.ignore()
            return
        path = paths[0]
        if path.suffix.casefold() not in self.extensions:
            permitidas = ", ".join(sorted(self.extensions))
            self.dropRejected.emit(
                f"Tipo de archivo no válido. Se esperaba: {permitidas}."
            )
            event.ignore()
            return
        self.set_file_path(str(path))
        event.acceptProposedAction()


class SalarHistoryDialog(QDialog):
    """Carga opcional del archivo histórico sin modificar el SALAR actual."""

    def __init__(self, current_path: str = "", parent=None):
        super().__init__(parent)
        self.setWindowTitle("SALAR histórico")
        self.setMinimumWidth(480)
        layout = QVBoxLayout(self)
        description = QLabel(
            "Selecciona un archivo SALAR anterior para poder identificar si existe alg\u00fan usuario "
            "que fue dado de baja en SALAR pero contin\u00faa habilitado en Dominio."
        )
        description.setWordWrap(True)
        layout.addWidget(description)

        label = QLabel("Archivo SALAR histórico")
        label.setObjectName("fileName")
        layout.addWidget(label)
        row = QHBoxLayout()
        self.file_edit = DropPathEdit({".txt"})
        self.file_edit.set_file_path(current_path)
        row.addWidget(self.file_edit, 1)
        browse = QPushButton("Buscar")
        browse.setObjectName("browseButton")
        browse.clicked.connect(self._browse)
        row.addWidget(browse)
        clear = QPushButton()
        clear.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_TrashIcon))
        clear.setObjectName("clearButton")
        clear.setToolTip("Quitar SALAR histórico")
        clear.clicked.connect(lambda: self.file_edit.set_file_path(""))
        row.addWidget(clear)
        layout.addLayout(row)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("Guardar selección")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _browse(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar SALAR histórico", "", TXT_FILTER
        )
        if path:
            self.file_edit.set_file_path(path)


class PathHintLabel(QLabel):
    """Muestra una ruta compacta y conserva la ruta completa en el tooltip."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self._path = ""
        self.setObjectName("pathHint")
        self.setMinimumWidth(0)
        self.setMaximumHeight(14)
        self.setVisible(False)

    def set_path(self, path: str):
        self._path = path or ""
        self.setToolTip(self._path)
        self.setAccessibleDescription(self._path)
        self.setVisible(bool(self._path))
        self._refresh_text()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._refresh_text()

    def _refresh_text(self):
        if self._path and self.width() > 0:
            metrics = QFontMetrics(self.font())
            self.setText(
                metrics.elidedText(
                    self._path, Qt.TextElideMode.ElideMiddle, self.width()
                )
            )
        else:
            self.clear()


class ConfigDialog(QDialog):
    """Editor de configuración de un sistema; muestra los valores actuales."""

    def __init__(self, config: SistemaConfig, parent=None):
        super().__init__(parent)
        self.config = config
        self.setWindowTitle(f"Configuración de {config.nombre}")
        self.setMinimumWidth(430)
        form = QFormLayout(self)
        self.inputs: dict[str, QLineEdit] = {}
        labels = {
            "columna_matriz": "Columna de acceso en matriz",
            "encabezado_matriz": "Encabezado de esa columna",
            "campo_login": "Campo del login en TXT",
            "campo_roles": "Campo de roles en TXT",
            "campos_txt": "Cantidad esperada de campos TXT",
            "separador_roles": "Separador de roles en TXT",
            "prefijo_login": "Prefijo a quitar del login",
            "sufijo_login": "Sufijo a quitar del login",
            "columna_matriz_organigrama": "Columna de organigrama en matriz",
            "encabezado_matriz_organigrama": "Encabezado de organigrama",
            "campo_organigrama": "Campo de organigrama en TXT",
            "separador_organigrama_matriz": "Separador en matriz",
            "separador_organigrama_txt": "Separador en TXT",
            "campo_estado": "Campo de estado en TXT",
            "estado_activo": "Valor de estado activo",
            "estado_inactivo": "Estado que se excluye",
        }
        if config.clave_cruce == "correo":
            labels["campo_login"] = "Campo del correo en TXT"
        visible = [
            "columna_matriz",
            "encabezado_matriz",
            "campo_login",
            "campo_roles",
            "campos_txt",
            "separador_roles",
            "prefijo_login",
            "sufijo_login",
        ]
        if config.columna_matriz_organigrama:
            visible += [
                "columna_matriz_organigrama",
                "encabezado_matriz_organigrama",
                "campo_organigrama",
                "separador_organigrama_matriz",
                "separador_organigrama_txt",
            ]
        if config.campo_estado:
            visible += ["campo_estado"]
            visible += [
                "estado_inactivo" if config.estado_inactivo else "estado_activo"
            ]
        for name in visible:
            value = getattr(config, name)
            edit = QLineEdit("" if value is None else str(value))
            edit.setPlaceholderText(
                "Sin ajuste" if name in ("prefijo_login", "sufijo_login") else ""
            )
            self.inputs[name] = edit
            form.addRow(labels[name], edit)
        self.null_checkbox = QCheckBox("Interpretar Null como ausencia de roles")
        self.null_checkbox.setChecked(config.null_sin_roles)
        form.addRow(self.null_checkbox)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save
            | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def _accept(self):
        values = {}
        integer_fields = {
            "campo_login",
            "campo_roles",
            "campos_txt",
            "campo_organigrama",
            "campo_estado",
        }
        for name, edit in self.inputs.items():
            value = edit.text().strip()
            if name in integer_fields:
                if not value.isdigit() or int(value) < 1:
                    QMessageBox.warning(
                        self,
                        "Configuración inválida",
                        f"{edit.text()!r} debe ser un número de campo mayor que cero.",
                    )
                    edit.setFocus()
                    return
                value = int(value)
            elif name in {"prefijo_login", "sufijo_login"}:
                value = value.upper()
            elif name.startswith("columna_matriz") and not value.isalpha():
                QMessageBox.warning(
                    self,
                    "Configuración inválida",
                    "La columna de matriz debe contener letras, por ejemplo AA.",
                )
                edit.setFocus()
                return
            elif not value and name not in {"prefijo_login", "sufijo_login"}:
                QMessageBox.warning(
                    self,
                    "Configuración incompleta",
                    f"Completa el campo: {name.replace('_', ' ')}.",
                )
                edit.setFocus()
                return
            values[name] = value
        values["null_sin_roles"] = self.null_checkbox.isChecked()
        self.config = replace(self.config, **values)
        self.accept()


class SourceConfigDialog(QDialog):
    """Editor de columnas para la matriz y las fuentes de personas."""

    def __init__(self, title: str, config, parent=None):
        super().__init__(parent)
        self.config = config
        self.setWindowTitle(f"Configuración de {title}")
        self.setMinimumWidth(430)
        form = QFormLayout(self)
        self.inputs: dict[str, QLineEdit] = {}
        labels = {
            "hoja": "Nombre de la hoja",
            "columna_cargo": "Columna del cargo",
            "campos_txt": "Cantidad de campos en TXT",
            "campos_nombre": "Campos del nombre (separados por coma)",
            "campo_cargo": "Campo del cargo",
            "campo_login": "Campo del login",
            "campo_oficina": "Campo de oficina (opcional)",
            "campo_gerencia": "Campo de gerencia (opcional)",
            "campo_correo": "Campo de correo (opcional)",
            "campo_carnet": "Campo del carnet (opcional)",
            "campo_telefono": "Campo del teléfono",
            "fila_encabezado": "Fila del encabezado (0 si no existe)",
            "campo_nombre": "Columna del nombre",
            "campo_cargo_actual": "Columna de Cargo Actual",
            "campo_cargo_designado": "Columna de Cargo Designado",
            "campo_oficina_actual": "Columna de Oficina Actual",
            "campo_oficina_designada": "Columna de Oficina Designada",
            "campo_fecha_inicio": "Columna de Fecha de inicio",
            "campo_fecha_fin": "Columna de Fecha de fin",
            "campo_correo": "Campo del correo en TXT",
            "campo_rol": "Campo del rol en TXT",
            "campo_estado": "Campo del estado en TXT (opcional)",
            "estado_inactivo": "Estado que se excluye",
            "hoja_codigos": "Hoja de códigos",
            "seccion_codigos": "Sección de códigos",
            "columna_seccion_codigos": "Columna de la sección",
            "columna_rol_matriz": "Columna Rol Matriz",
            "columna_codigo_txt": "Columna de rol/código TXT",
            "columna_matriz": "Columna de matriz del sistema",
        }
        if isinstance(config, SistemaExternoConfig):
            labels["campo_nombre"] = "Campo del nombre en TXT (opcional)"
        for field in fields(config):
            if field.name == "nombre_sistema":
                continue
            value = getattr(config, field.name)
            if isinstance(value, tuple):
                value = ", ".join(map(str, value))
            edit = QLineEdit("" if value is None else str(value))
            self.inputs[field.name] = edit
            form.addRow(labels[field.name], edit)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save
            | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def _accept(self):
        values = {}
        integer_fields = {
            "campos_txt",
            "campo_cargo",
            "campo_login",
            "campo_oficina",
            "campo_gerencia",
            "campo_correo",
            "campo_carnet",
            "fila_encabezado",
            "campo_nombre",
            "campo_estado",
            "campo_cargo_actual",
            "campo_cargo_designado",
            "campo_rol",
            "campo_oficina_actual",
            "campo_oficina_designada",
            "campo_fecha_inicio",
            "campo_fecha_fin",
            "campo_telefono",
        }
        for name, edit in self.inputs.items():
            value = edit.text().strip()
            if name in integer_fields:
                if (
                    isinstance(self.config, PersonasConfig)
                    and name
                    in {
                        "campo_oficina",
                        "campo_gerencia",
                        "campo_correo",
                        "campo_carnet",
                    }
                    and not value
                ):
                    values[name] = None
                    continue
                if (isinstance(self.config, SistemaExternoConfig)
                        and name in {"campo_nombre", "campo_estado"}
                        and not value):
                    values[name] = None
                    continue
                if not value.isdigit() or (
                    int(value) < 0 if name == "fila_encabezado" else int(value) < 1
                ):
                    QMessageBox.warning(
                        self,
                        "Configuración inválida",
                        f"{edit.text()!r} no es una posición válida para {name.replace('_', ' ')}.",
                    )
                    edit.setFocus()
                    return
                value = int(value)
            elif name == "campos_nombre":
                try:
                    value = tuple(
                        int(parte.strip())
                        for parte in value.split(",")
                        if parte.strip()
                    )
                except ValueError:
                    value = ()
                if not value or any(posicion < 1 for posicion in value):
                    QMessageBox.warning(
                        self,
                        "Configuración inválida",
                        "Indica campos positivos separados por comas, por ejemplo 1, 2, 3.",
                    )
                    edit.setFocus()
                    return
            elif name in {
                "columna_cargo",
                "columna_seccion_codigos",
                "columna_rol_matriz",
                "columna_codigo_txt",
                "columna_matriz",
            }:
                value = value.upper()
                if value.isalpha():
                    values[name] = value
                    continue
                QMessageBox.warning(
                    self,
                    "Configuración inválida",
                    "La columna debe contener letras, por ejemplo B o AA.",
                )
                edit.setFocus()
                return
            elif (name == "estado_inactivo"
                    and isinstance(self.config, SistemaExternoConfig)
                    and values.get("campo_estado") is None):
                values[name] = ""
                continue
            elif not value:
                QMessageBox.warning(
                    self,
                    "Configuración incompleta",
                    f"Completa el campo: {name.replace('_', ' ')}.",
                )
                edit.setFocus()
                return
            values[name] = value
        if isinstance(self.config, PersonasConfig) and any(
            posicion > values["campos_txt"]
            for posicion in (
                *values["campos_nombre"],
                values["campo_cargo"],
                values["campo_login"],
                *(
                    p
                    for p in (
                        values["campo_oficina"],
                        values["campo_gerencia"],
                        values["campo_correo"],
                        values["campo_carnet"],
                    )
                    if p is not None
                ),
            )
        ):
            QMessageBox.warning(
                self,
                "Configuración inválida",
                "Los campos configurados deben caber dentro del TXT.",
            )
            return
        if isinstance(self.config, SistemaExternoConfig) and any(
            value is not None and value > values["campos_txt"]
            for value in (values["campo_correo"], values["campo_nombre"],
                          values["campo_rol"], values["campo_estado"])
        ):
            QMessageBox.warning(
                self,
                "Configuración inválida",
                "Los campos configurados deben caber dentro del TXT.",
            )
            return
        if (isinstance(self.config, SistemaExternoConfig)
                and values["campo_estado"] is not None
                and not values["estado_inactivo"].strip()):
            QMessageBox.warning(
                self,
                "Configuración inválida",
                "Indica qué estado debe excluirse del análisis.",
            )
            return
        self.config = replace(self.config, **values)
        self.accept()


# Los workers mantienen las tareas de lectura y generación fuera del hilo visual.
class AnalysisWorker(QThread):
    completed = Signal(str, int, int)
    failed = Signal(str)

    def __init__(
        self,
        matrix: Path,
        people_file: Path,
        source: str,
        files: dict[str, Path],
        output: Path,
        configs: tuple[SistemaConfig, ...],
        temporales: Path | None = None,
        matrix_config: MatrizConfig = MatrizConfig(),
        people_config: PersonasConfig = SALAR_CONFIG,
        temporales_config: TemporalesConfig = TemporalesConfig(),
        domain_file: Path | None = None,
        domain_config: PersonasConfig = DOMINIO_CONFIG,
        asfi_file: Path | None = None,
        asfi_config: ASFIConfig = ASFIConfig(),
        salar_anterior_file: Path | None = None,
        infocred_file: Path | None = None,
        infocred_config: INFOCREDConfig = INFOCREDConfig(),
    ):
        super().__init__()
        self.arguments = (matrix, people_file, files, output, configs, temporales)
        self.output_file = output
        self.source = source
        self.matrix_config = matrix_config
        self.people_config = people_config
        self.temporales_config = temporales_config
        self.domain_file = domain_file
        self.domain_config = domain_config
        self.asfi_file = asfi_file
        self.asfi_config = asfi_config
        self.salar_anterior_file = salar_anterior_file
        self.infocred_file = infocred_file
        self.infocred_config = infocred_config

    def run(self):
        # Ejecuta el análisis completo y notifica a la ventana por señales Qt.
        try:
            output, result = ejecutar_revision_completa(
                *self.arguments,
                fuente_personas=self.source,
                config_matriz=self.matrix_config,
                config_personas=self.people_config,
                config_temporales=self.temporales_config,
                archivo_dominio=self.domain_file,
                config_dominio=self.domain_config,
                archivo_asfi=self.asfi_file,
                config_asfi=self.asfi_config,
                archivo_salar_anterior=self.salar_anterior_file,
                archivo_infocred=self.infocred_file,
                config_infocred=self.infocred_config,
            )
            self.completed.emit(
                str(output),
                result.get("observaciones_totales", len(result["observaciones"])),
                len(result["errores"]),
            )
        except Exception as exc:
            self.output_file.unlink(missing_ok=True)
            self.failed.emit(str(exc))


class ReportWorker(QThread):
    completed = Signal(str)
    failed = Signal(str)

    def __init__(
        self,
        analysis_file: Path,
        output_file: Path,
        false_positives_file: Path | None = None,
    ):
        super().__init__()
        self.analysis_file = analysis_file
        self.output_file = output_file
        self.false_positives_file = false_positives_file

    def run(self):
        # Convierte el análisis guardado en el informe basado en plantilla.
        try:
            path = generar_informe(
                self.analysis_file,
                DEFAULT_TEMPLATE,
                self.output_file,
                falsos_positivos=self.false_positives_file,
            )
            self.completed.emit(str(path))
        except Exception as exc:
            self.output_file.unlink(missing_ok=True)
            self.failed.emit(str(exc))


class TelephonyWorker(QThread):
    completed = Signal(str, int)
    failed = Signal(str)

    def __init__(
        self,
        domain_file: Path,
        internal_file: Path,
        output_file: Path,
        domain_config: TelefoniaConfig,
        isabel_config: TelefoniaConfig,
    ):
        super().__init__()
        self.domain_file = domain_file
        self.internal_file = internal_file
        self.output_file = output_file
        self.domain_config = domain_config
        self.isabel_config = isabel_config

    def run(self):
        # Genera el reporte de telefonía en segundo plano.
        try:
            output, summary = generar_reporte_telefonia(
                self.domain_file,
                self.internal_file,
                self.output_file,
                self.domain_config,
                self.isabel_config,
            )
            self.completed.emit(str(output), summary["observaciones"])
        except Exception as exc:
            self.output_file.unlink(missing_ok=True)
            self.failed.emit(str(exc))


class PdfExcelWorker(QThread):
    completed = Signal(str, int)
    failed = Signal(str)

    def __init__(self, pdf_file: Path):
        super().__init__()
        self.pdf_file = pdf_file

    def run(self):
        # Genera el Excel temporal sin bloquear la ventana principal.
        temporal = None
        try:
            temporal = _crear_ruta_temporal_excel("roles_asfi_")
            output = convertir_pdf_a_excel(self.pdf_file, Path(temporal))
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                registros = max(0, workbook["Roles ASFI"].max_row - 1)
            finally:
                workbook.close()
            self.completed.emit(str(output), registros)
        except Exception as exc:
            if temporal:
                Path(temporal).unlink(missing_ok=True)
            self.failed.emit(str(exc))


class SintesisExcelWorker(QThread):
    completed = Signal(str, int)
    failed = Signal(str)

    def __init__(self, pdf_file: Path):
        super().__init__()
        self.pdf_file = pdf_file

    def run(self):
        # Conserva el parser de Síntesis y genera el libro fuera del hilo visual.
        temporal = None
        try:
            temporal = _crear_ruta_temporal_excel("sintesis_")
            output, registros = convertir_pdf_sintesis_a_excel(self.pdf_file, temporal)
            self.completed.emit(str(output), registros)
        except Exception as exc:
            if temporal:
                temporal.unlink(missing_ok=True)
            self.failed.emit(str(exc))


class MainWindow(QMainWindow):
    def __init__(self):
        # Carga ajustes guardados, crea las pantallas y conecta las acciones.
        super().__init__()
        self.setWindowTitle("Analizador de Accesos")
        icon_path = recurso_aplicacion("assets/icono.ico")
        if icon_path.is_file():
            self.setWindowIcon(QIcon(str(icon_path)))
        self.resize(1360, 880)
        self.setMinimumSize(760, 600)
        self.settings = QSettings("RevisionAccesos", "AnalizadorAccesos")
        self.configs = {config.nombre: self._load_config(config) for config in SISTEMAS}
        self.source_configs = {
            "matriz": self._load_source_config("matriz", MatrizConfig()),
            "dominio": self._load_source_config("dominio", DOMINIO_CONFIG),
            "salar": self._load_source_config("salar", SALAR_CONFIG),
            "temporales": self._load_source_config("temporales", TemporalesConfig()),
            "asfi": self._load_source_config("asfi", ASFIConfig()),
            "infocred": self._load_source_config("infocred", INFOCREDConfig()),
            "telefonia_dominio": self._load_source_config(
                "telefonia_dominio", TelefoniaConfig(1, 11)
            ),
            "telefonia_isabel": self._load_source_config(
                "telefonia_isabel", TelefoniaConfig(1, 2)
            ),
        }
        self.system_rows: dict[str, DropPathEdit] = {}
        self.salar_anterior_path = ""
        self.worker: AnalysisWorker | None = None
        self.report_worker: ReportWorker | None = None
        self.telephony_worker: TelephonyWorker | None = None
        self.pdf_excel_worker: PdfExcelWorker | None = None
        self.sintesis_excel_worker: SintesisExcelWorker | None = None
        self._build_ui()

    # Lee y guarda ajustes del usuario para conservar la configuración entre sesiones.
    def _load_config(self, default: SistemaConfig) -> SistemaConfig:
        self.settings.beginGroup(f"sistemas/{default.nombre}")
        changes = {}
        for field in fields(default):
            if self.settings.contains(field.name):
                current = getattr(default, field.name)
                value = self.settings.value(field.name)
                if isinstance(current, bool):
                    value = str(value).lower() in {"true", "1"}
                elif isinstance(current, int):
                    value = int(value)
                elif (
                    current is None
                    and field.name
                    in {
                        "campo_organigrama",
                        "campo_estado",
                        "campo_nombre",
                        "campo_cargo",
                        "campo_oficina",
                        "campo_gerencia",
                        "campo_login_respaldo",
                    }
                    and value
                ):
                    value = int(value)
                elif current is None and not value:
                    value = None
                changes[field.name] = value
        self.settings.endGroup()
        return replace(default, **changes)

    def _save_config(self, config: SistemaConfig):
        self.settings.beginGroup(f"sistemas/{config.nombre}")
        for field in fields(config):
            value = getattr(config, field.name)
            self.settings.setValue(field.name, "" if value is None else value)
        self.settings.endGroup()
        self.settings.sync()

    def _load_source_config(self, name: str, default):
        self.settings.beginGroup(f"fuentes/{name}")
        changes = {}
        for field in fields(default):
            if not self.settings.contains(field.name):
                continue
            value = self.settings.value(field.name)
            current = getattr(default, field.name)
            if isinstance(current, tuple):
                if isinstance(value, str):
                    value = tuple(
                        int(part.strip()) for part in value.split(",") if part.strip()
                    )
                else:
                    value = tuple(int(part) for part in value)
            elif (field.name in {
                "campo_oficina", "campo_gerencia", "campo_correo", "campo_carnet",
                "campo_estado",
            } or (field.name == "campo_nombre" and current is None)):
                value = (
                    int(value)
                    if str(value or "").strip()
                    else getattr(default, field.name)
                    if name == "dominio"
                    else None
                )
            elif isinstance(current, int):
                value = int(value)
            else:
                value = str(value)
            changes[field.name] = value
        self.settings.endGroup()
        return replace(default, **changes)

    def _save_source_config(self, name: str, config):
        self.settings.beginGroup(f"fuentes/{name}")
        for field in fields(config):
            value = getattr(config, field.name)
            if isinstance(value, tuple):
                value = ",".join(map(str, value))
            elif value is None:
                value = ""
            self.settings.setValue(field.name, value)
        self.settings.endGroup()
        self.settings.sync()

    # Construye las páginas, controles de carga y botones de análisis.
    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QHBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        self.sidebar = QFrame()
        self.sidebar.setObjectName("sidebar")
        side_layout = QVBoxLayout(self.sidebar)
        self.sidebar_layout = side_layout
        side_layout.setContentsMargins(14, 22, 14, 18)
        side_layout.setSpacing(12)
        self.brand_title = QLabel("BDP")
        self.brand_title.setObjectName("brandTitle")
        self.brand_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        side_layout.addWidget(self.brand_title)
        self.brand_logo_card = QFrame()
        self.brand_logo_card.setObjectName("brandLogoCard")
        self.brand_logo_card.setFixedHeight(72)
        logo_layout = QHBoxLayout(self.brand_logo_card)
        logo_layout.setContentsMargins(3, 3, 3, 3)
        self.brand_logo = QLabel()
        self.brand_logo.setObjectName("brandLogo")
        self.brand_logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        source_logo = QPixmap(str(recurso_aplicacion("assets/logo_bdp.jpeg")))
        if not source_logo.isNull():
            logo_crop = source_logo.copy(18, 135, 411, 157)
            self.brand_logo.setPixmap(
                logo_crop.scaled(
                    152,
                    62,
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
            )
        else:
            self.brand_logo.setText("BDP")
        logo_layout.addWidget(self.brand_logo)
        side_layout.addWidget(self.brand_logo_card)
        self.brand_title.setVisible(False)
        self.nav_analysis_button = QPushButton("◉   Analizar sistemas")
        self.nav_analysis_button.setObjectName("navButton")
        self.nav_analysis_button.setCheckable(True)
        self.nav_analysis_button.setChecked(True)
        self.nav_analysis_button.setMinimumHeight(42)
        self.nav_analysis_button.clicked.connect(lambda: self._show_view(0))
        side_layout.addWidget(self.nav_analysis_button)
        self.nav_report_button = QPushButton("▤   Generar Reporte")
        self.nav_report_button.setObjectName("navButton")
        self.nav_report_button.setCheckable(True)
        self.nav_report_button.setMinimumHeight(42)
        self.nav_report_button.clicked.connect(lambda: self._show_view(1))
        side_layout.addWidget(self.nav_report_button)
        self.nav_pdf_excel_button = QPushButton(PDF_EXCEL_NAV_TEXT)
        self.nav_pdf_excel_button.setObjectName("navButton")
        self.nav_pdf_excel_button.setCheckable(True)
        self.nav_pdf_excel_button.setMinimumHeight(54)
        self.nav_pdf_excel_button.setToolTip("Conversor PDF a Excel")
        self.nav_pdf_excel_button.clicked.connect(lambda: self._show_view(3))
        side_layout.addWidget(self.nav_pdf_excel_button)
        self.nav_telephony_button = QPushButton("☎   Telefonía")
        self.nav_telephony_button.setObjectName("navButton")
        self.nav_telephony_button.setCheckable(True)
        self.nav_telephony_button.setMinimumHeight(42)
        self.nav_telephony_button.clicked.connect(lambda: self._show_view(2))
        side_layout.addWidget(self.nav_telephony_button)
        side_layout.addStretch(1)
        self.sidebar_footer = QLabel(
            "Seguridad de la Información\nGerencia de Sistemas y\nTecnología de la Información"
        )
        self.sidebar_footer.setObjectName("sidebarFooter")
        self.sidebar_footer.setWordWrap(True)
        side_layout.addWidget(self.sidebar_footer)
        root.addWidget(self.sidebar)

        content = QWidget()
        content.setObjectName("content")
        outer = QVBoxLayout(content)
        outer.setContentsMargins(24, 18, 24, 16)
        outer.setSpacing(12)

        header = QHBoxLayout()
        title_block = QVBoxLayout()
        title_block.setSpacing(2)
        eyebrow = QLabel("CONTROL DE ACCESOS  /  BDP")
        eyebrow.setObjectName("eyebrow")
        self.page_title = QLabel("Analizador de Accesos")
        self.page_title.setObjectName("title")
        title_block.addWidget(eyebrow)
        title_block.addWidget(self.page_title)
        header.addLayout(title_block, 1)
        badge = QLabel("REVISIÓN DE SISTEMAS INTERNOS Y EXTERNOS")
        badge.setObjectName("versionBadge")
        badge.setAlignment(Qt.AlignmentFlag.AlignCenter)
        header.addWidget(badge, 0, Qt.AlignmentFlag.AlignTop)
        outer.addLayout(header)

        scroll = QScrollArea()
        scroll.setObjectName("contentScroll")
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        page = QWidget()
        page.setObjectName("contentPage")
        page_layout = QVBoxLayout(page)
        page_layout.setContentsMargins(0, 2, 0, 8)
        page_layout.setSpacing(14)
        self.cards_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.cards_splitter.setObjectName("cardsSplitter")
        self.cards_splitter.setChildrenCollapsible(False)

        primary_panel = QWidget()
        primary_layout = QVBoxLayout(primary_panel)
        primary_layout.setContentsMargins(0, 0, 0, 0)
        primary_layout.setSpacing(12)
        files_card = self._card("ARCHIVOS PRINCIPALES")
        files_note = QLabel(
            "La revisión usa SALAR como fuente principal y Dominio como fuente complementaria. "
            "Carga cada archivo con Buscar o arrástralo a su recuadro."
        )
        files_note.setObjectName("hint")
        files_card.layout().addWidget(files_note)
        self.matrix_edit = self._file_row(
            files_card,
            "Matriz de\naccesos",
            MATRIX_FILTER,
            {".xlsx", ".xlsm"},
            "EXCEL · .xlsx / .xlsm",
            lambda: self._configure_source("matriz", "Matriz"),
            label_width=112,
        )
        self.domain_edit = self._file_row(
            files_card,
            "Dominio",
            TXT_FILTER,
            {".txt"},
            "TEXTO · .txt",
            lambda: self._configure_source("dominio", "Dominio"),
            label_width=112,
        )
        self.salar_edit = self._file_row(
            files_card,
            "Salar",
            TXT_FILTER,
            {".txt"},
            "TEXTO · .txt",
            lambda: self._configure_source("salar", "SALAR"),
            label_width=112,
            extra_action=self._configure_salar_history,
        )
        self.temporales_edit = self._file_row(
            files_card,
            "Accesos temporales (opcional)",
            MATRIX_FILTER,
            {".xlsx", ".xlsm"},
            "EXCEL · .xlsx / .xlsm",
            lambda: self._configure_source("temporales", "Accesos temporales"),
            label_width=112,
        )
        primary_layout.addWidget(files_card)

        output_card = QFrame()
        output_card.setObjectName("actionCard")
        output_layout = QVBoxLayout(output_card)
        output_layout.setContentsMargins(16, 12, 16, 12)
        output_layout.setSpacing(8)
        output_label = QLabel("REVISIÓN DE ACCESOS")
        output_label.setObjectName("sectionTitle")
        output_layout.addWidget(output_label)
        action_row = QHBoxLayout()
        action_row.setSpacing(8)
        self.clear_button = QPushButton("Limpiar archivos")
        self.clear_button.setObjectName("secondaryButton")
        self.clear_button.setMinimumHeight(46)
        self.clear_button.clicked.connect(self._clear_selection)
        action_row.addWidget(self.clear_button, 0, Qt.AlignmentFlag.AlignTop)
        analyze_action = QVBoxLayout()
        analyze_action.setContentsMargins(0, 0, 0, 0)
        analyze_action.setSpacing(16)
        self.analyze_salar_button = QPushButton("▶   Generar Análisis")
        self.analyze_salar_button.setObjectName("analyzeButton")
        self.analyze_salar_button.setMinimumHeight(46)
        self.analyze_salar_button.clicked.connect(self._analyze)
        analyze_action.addWidget(self.analyze_salar_button)
        output_note = QLabel(
            "Cuando el análisis esté listo, podrás elegir dónde guardarlo."
        )
        output_note.setObjectName("hint")
        analyze_action.addWidget(output_note)
        action_row.addLayout(analyze_action, 1)
        output_layout.addLayout(action_row)
        primary_layout.addWidget(output_card)

        report_card = QFrame()
        report_card.setObjectName("actionCard")
        report_layout = QVBoxLayout(report_card)
        report_layout.setContentsMargins(16, 12, 16, 12)
        report_layout.setSpacing(8)
        report_title = QLabel(
            "GENERAR REPORTE DE REVISIÓN DE SISTEMAS INTERNOS/EXTERNOS"
        )
        report_title.setObjectName("sectionTitle")
        report_title.setWordWrap(True)
        report_layout.addWidget(report_title)
        report_note = QLabel(
            "Carga la revisión de accesos y, si tienes, el Excel de falsos positivos."
        )
        report_note.setObjectName("hint")
        report_note.setWordWrap(True)
        report_layout.addWidget(report_note)
        report_inputs = QHBoxLayout()
        report_inputs.setSpacing(10)
        analysis_box, self.report_analysis_edit = self._report_file_box(
            "Revisión de accesos", required=True
        )
        false_positive_box, self.false_positives_edit = self._report_file_box(
            "Falsos positivos (opcional)", required=False
        )
        report_inputs.addWidget(analysis_box, 1)
        report_inputs.addWidget(false_positive_box, 1)
        report_layout.addLayout(report_inputs)

        report_action_row = QHBoxLayout()
        report_action_row.setSpacing(8)
        self.clear_report_button = QPushButton("Limpiar archivos")
        self.clear_report_button.setObjectName("secondaryButton")
        self.clear_report_button.setMinimumHeight(44)
        self.clear_report_button.clicked.connect(self._clear_report_selection)
        report_action_row.addWidget(
            self.clear_report_button, 0, Qt.AlignmentFlag.AlignTop
        )
        report_action = QVBoxLayout()
        report_action.setContentsMargins(0, 0, 0, 0)
        report_action.setSpacing(16)
        self.generate_report_button = QPushButton("▣   Generar Reporte")
        self.generate_report_button.setObjectName("analyzeButton")
        self.generate_report_button.setMinimumHeight(44)
        self.generate_report_button.clicked.connect(self._generate_report)
        report_action.addWidget(self.generate_report_button)
        report_save_note = QLabel(
            "Cuando el reporte esté listo, podrás elegir dónde guardarlo."
        )
        report_save_note.setObjectName("hint")
        report_action.addWidget(report_save_note)
        report_action_row.addLayout(report_action, 1)
        report_layout.addLayout(report_action_row)

        report_page = QWidget()
        report_page.setObjectName("contentPage")
        report_page_layout = QVBoxLayout(report_page)
        report_page_layout.setContentsMargins(0, 2, 0, 8)
        report_page_layout.addWidget(report_card, 0, Qt.AlignmentFlag.AlignTop)
        report_page_layout.addStretch(1)
        report_scroll = QScrollArea()
        report_scroll.setObjectName("contentScroll")
        report_scroll.setWidgetResizable(True)
        report_scroll.setFrameShape(QFrame.Shape.NoFrame)
        report_scroll.setWidget(report_page)

        telephony_page = QWidget()
        telephony_page.setObjectName("contentPage")
        telephony_page_layout = QVBoxLayout(telephony_page)
        telephony_page_layout.setContentsMargins(0, 2, 0, 8)
        telephony_page_layout.setSpacing(14)
        telephony_card = self._card("SISTEMA DE TELEFONÍA")
        telephony_note = QLabel(
            "Internos (Isabel) es la fuente de referencia. Se agrupan sus teléfonos por persona y "
            "se comparan con Dominio; los nombres aproximados quedan para validar."
        )
        telephony_note.setObjectName("hint")
        telephony_note.setWordWrap(True)
        telephony_card.layout().addWidget(telephony_note)
        self.telephony_domain_edit = self._file_row(
            telephony_card,
            "Dominio",
            TXT_FILTER,
            {".txt"},
            "TXT",
            lambda: self._configure_source("telefonia_dominio", "Dominio"),
            label_width=150,
        )
        self.telephony_internal_edit = self._file_row(
            telephony_card,
            "Internos (Isabel)",
            TXT_FILTER,
            {".txt"},
            "TXT",
            lambda: self._configure_source("telefonia_isabel", "Internos (Isabel)"),
            label_width=150,
        )

        telephony_actions = QHBoxLayout()
        telephony_actions.setSpacing(8)
        self.clear_telephony_button = QPushButton("Limpiar archivos")
        self.clear_telephony_button.setObjectName("secondaryButton")
        self.clear_telephony_button.setMinimumHeight(46)
        self.clear_telephony_button.clicked.connect(self._clear_telephony_selection)
        telephony_actions.addWidget(
            self.clear_telephony_button, 0, Qt.AlignmentFlag.AlignTop
        )
        telephony_action = QVBoxLayout()
        telephony_action.setContentsMargins(0, 0, 0, 0)
        telephony_action.setSpacing(16)
        self.generate_telephony_button = QPushButton("◉   Generar Análisis")
        self.generate_telephony_button.setObjectName("analyzeButton")
        self.generate_telephony_button.setMinimumHeight(46)
        self.generate_telephony_button.clicked.connect(self._generate_telephony_report)
        telephony_action.addWidget(self.generate_telephony_button)
        telephony_save_note = QLabel(
            "Cuando el análisis de telefonía esté listo, podrás elegir dónde guardarlo."
        )
        telephony_save_note.setObjectName("hint")
        telephony_action.addWidget(telephony_save_note)
        telephony_actions.addLayout(telephony_action, 1)
        telephony_card.layout().addLayout(telephony_actions)
        telephony_page_layout.addWidget(telephony_card, 0, Qt.AlignmentFlag.AlignTop)
        telephony_page_layout.addStretch(1)
        telephony_scroll = QScrollArea()
        telephony_scroll.setObjectName("contentScroll")
        telephony_scroll.setWidgetResizable(True)
        telephony_scroll.setFrameShape(QFrame.Shape.NoFrame)
        telephony_scroll.setWidget(telephony_page)

        pdf_excel_page = QWidget()
        pdf_excel_page.setObjectName("contentPage")
        pdf_excel_layout = QVBoxLayout(pdf_excel_page)
        pdf_excel_layout.setContentsMargins(12, 12, 12, 12)
        pdf_excel_layout.setSpacing(14)
        pdf_excel_card = self._card("CONVERSOR PDF ASFI A EXCEL")
        pdf_excel_card.setMaximumWidth(760)
        pdf_excel_card.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred
        )
        pdf_excel_note = QLabel(
            "Selecciona el reporte PDF de usuarios y roles vigentes de ASFI. "
            "Se extraerán el correo, el nombre y una fila por cada rol; "
            "los usuarios sin roles también se conservarán."
        )
        pdf_excel_note.setObjectName("hint")
        pdf_excel_note.setWordWrap(True)
        pdf_excel_card.layout().addWidget(pdf_excel_note)
        self.pdf_excel_input = self._file_row(
            pdf_excel_card,
            "Reporte ASFI (PDF)",
            "Archivos PDF (*.pdf)",
            {".pdf"},
            "PDF · .pdf",
            label_width=150,
        )

        pdf_excel_actions = QHBoxLayout()
        pdf_excel_actions.setContentsMargins(0, 0, 0, 0)
        pdf_excel_actions.setSpacing(8)
        self.clear_pdf_excel_button = QPushButton("Limpiar archivos")
        self.clear_pdf_excel_button.setObjectName("secondaryButton")
        self.clear_pdf_excel_button.setMinimumHeight(46)
        self.clear_pdf_excel_button.clicked.connect(self._clear_pdf_excel_selection)
        pdf_excel_actions.addWidget(self.clear_pdf_excel_button)
        self.convert_pdf_excel_button = QPushButton("Generar")
        self.convert_pdf_excel_button.setObjectName("analyzeButton")
        self.convert_pdf_excel_button.setMinimumHeight(46)
        self.convert_pdf_excel_button.clicked.connect(self._convert_pdf_to_excel)
        pdf_excel_actions.addWidget(self.convert_pdf_excel_button, 1)
        pdf_excel_card.layout().addLayout(pdf_excel_actions)
        pdf_excel_save_note = QLabel(
            "Cuando el Excel de Asfi este listo, podrás elegir dónde guardarlo."
        )
        pdf_excel_save_note.setObjectName("hint")
        pdf_excel_save_note.setWordWrap(True)
        pdf_excel_card.layout().addWidget(pdf_excel_save_note)
        pdf_excel_layout.addWidget(pdf_excel_card)

        sintesis_pdf_card = self._card("CONVERSOR PDF SÍNTESIS A EXCEL")
        sintesis_pdf_card.setMaximumWidth(760)
        sintesis_pdf_card.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred
        )
        sintesis_pdf_note = QLabel(
            "Selecciona el reporte PDF de operadores de Síntesis para extraer sus datos "
            "y generar el Excel correspondiente."
        )
        sintesis_pdf_note.setObjectName("hint")
        sintesis_pdf_note.setWordWrap(True)
        sintesis_pdf_card.layout().addWidget(sintesis_pdf_note)
        self.sintesis_pdf_input = self._file_row(
            sintesis_pdf_card,
            "Reporte Síntesis (PDF)",
            "Archivos PDF (*.pdf)",
            {".pdf"},
            "PDF · .pdf",
            label_width=150,
        )

        sintesis_actions = QHBoxLayout()
        sintesis_actions.setContentsMargins(0, 0, 0, 0)
        sintesis_actions.setSpacing(8)
        self.clear_sintesis_pdf_button = QPushButton("Limpiar archivos")
        self.clear_sintesis_pdf_button.setObjectName("secondaryButton")
        self.clear_sintesis_pdf_button.setMinimumHeight(46)
        self.clear_sintesis_pdf_button.clicked.connect(
            self._clear_sintesis_pdf_selection
        )
        sintesis_actions.addWidget(self.clear_sintesis_pdf_button)
        self.convert_sintesis_pdf_button = QPushButton("Generar")
        self.convert_sintesis_pdf_button.setObjectName("analyzeButton")
        self.convert_sintesis_pdf_button.setMinimumHeight(46)
        self.convert_sintesis_pdf_button.clicked.connect(
            self._convert_sintesis_pdf_to_excel
        )
        sintesis_actions.addWidget(self.convert_sintesis_pdf_button, 1)
        sintesis_pdf_card.layout().addLayout(sintesis_actions)
        sintesis_save_note = QLabel(
            "Cuando el Excel de Síntesis esté listo, podrás elegir dónde guardarlo."
        )
        sintesis_save_note.setObjectName("hint")
        sintesis_save_note.setWordWrap(True)
        sintesis_pdf_card.layout().addWidget(sintesis_save_note)
        pdf_excel_layout.addWidget(sintesis_pdf_card)
        pdf_excel_layout.addStretch(2)
        pdf_excel_scroll = QScrollArea()
        pdf_excel_scroll.setObjectName("contentScroll")
        pdf_excel_scroll.setWidgetResizable(True)
        pdf_excel_scroll.setFrameShape(QFrame.Shape.NoFrame)
        pdf_excel_scroll.setWidget(pdf_excel_page)

        primary_layout.addStretch(1)

        systems_card = QFrame()
        systems_card.setObjectName("systemsCard")
        systems_card_layout = QVBoxLayout(systems_card)
        systems_card_layout.setContentsMargins(12, 8, 12, 8)
        systems_card_layout.setSpacing(8)
        systems_scroll = QScrollArea()
        systems_scroll.setObjectName("systemRowsScroll")
        systems_scroll.setWidgetResizable(True)
        systems_scroll.setFrameShape(QFrame.Shape.NoFrame)
        systems_scroll.setMinimumHeight(200)
        systems_content = QWidget()
        systems_content.setObjectName("systemRowsContent")
        systems_layout = QVBoxLayout(systems_content)
        systems_layout.setContentsMargins(0, 0, 0, 0)
        systems_layout.setSpacing(5)
        internal_heading = QLabel("SISTEMAS INTERNOS")
        internal_heading.setObjectName("sectionTitle")
        internal_heading.setAlignment(Qt.AlignmentFlag.AlignCenter)
        systems_layout.addWidget(internal_heading)
        internal_note = QLabel(
            "Carga el TXT de cada sistema que quieras incluir en la revisión de accesos."
        )
        internal_note.setObjectName("hint")
        internal_note.setAlignment(Qt.AlignmentFlag.AlignCenter)
        internal_note.setWordWrap(True)
        systems_layout.addWidget(internal_note)
        self.system_count_label = QLabel()
        self.system_count_label.setObjectName("countBadge")
        self.system_count_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        systems_layout.addWidget(
            self.system_count_label, 0, Qt.AlignmentFlag.AlignCenter
        )
        for config in SISTEMAS:
            display_name = {
                "BDPSAM": "ArgosForm 2doPiso (BDPSAM)",
                "FG": "Fondo de Garantia (FG)",
            }.get(config.nombre, config.nombre)
            path_edit = self._file_row(
                systems_content,
                display_name,
                TXT_FILTER,
                {".txt"},
                "TXT",
                lambda _checked=False, name=config.nombre: self._configure(name),
                label_width=112,
            )
            path_edit.fileChanged.connect(
                lambda _path: self._actualizar_conteo_sistemas()
            )
            self.system_rows[config.nombre] = path_edit
        external_heading = QLabel("SISTEMAS EXTERNOS")
        external_heading.setObjectName("sectionTitle")
        external_heading.setContentsMargins(0, 10, 0, 0)
        external_heading.setAlignment(Qt.AlignmentFlag.AlignCenter)
        systems_layout.addWidget(external_heading)
        external_note = QLabel(
            "Puedes analizar solo ASFI o INFOCRED; no es obligatorio cargar un sistema interno."
        )
        external_note.setObjectName("hint")
        external_note.setAlignment(Qt.AlignmentFlag.AlignCenter)
        external_note.setWordWrap(True)
        systems_layout.addWidget(external_note)
        self.external_count_label = QLabel()
        self.external_count_label.setObjectName("countBadge")
        self.external_count_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        systems_layout.addWidget(
            self.external_count_label, 0, Qt.AlignmentFlag.AlignCenter
        )
        self.asfi_edit = self._file_row(
            systems_content,
            "ASFI",
            TXT_FILTER,
            {".txt"},
            "TXT",
            lambda: self._configure_source("asfi", "ASFI"),
            label_width=112,
        )
        self.external_system_rows = {"ASFI": self.asfi_edit}
        self.asfi_edit.fileChanged.connect(
            lambda _path: self._actualizar_conteo_sistemas()
        )
        self.infocred_edit = self._file_row(
            systems_content,
            "INFOCRED",
            TXT_FILTER,
            {".txt"},
            "TXT",
            lambda: self._configure_source("infocred", "INFOCRED"),
            label_width=112,
        )
        self.external_system_rows["INFOCRED"] = self.infocred_edit
        self.infocred_edit.fileChanged.connect(
            lambda _path: self._actualizar_conteo_sistemas()
        )
        systems_layout.addStretch(1)
        systems_scroll.setWidget(systems_content)
        systems_card_layout.addWidget(systems_scroll, 1)

        self.cards_splitter.addWidget(primary_panel)
        self.cards_splitter.addWidget(systems_card)
        # Let the splitter honor its intended proportions instead of being
        # forced to the children’s large size hints.
        primary_panel.setMinimumWidth(0)
        systems_card.setMinimumWidth(0)
        self.cards_splitter.setStretchFactor(0, 1)
        self.cards_splitter.setStretchFactor(1, 1)
        self._initial_split_applied = False
        page_layout.addWidget(self.cards_splitter)
        scroll.setWidget(page)
        self.main_scroll = scroll
        self.views = QStackedWidget()
        self.views.addWidget(scroll)
        self.views.addWidget(report_scroll)
        self.views.addWidget(telephony_scroll)
        self.views.addWidget(pdf_excel_scroll)
        outer.addWidget(self.views, 1)

        bottom = QHBoxLayout()
        self.status = QLabel("Listo para analizar.")
        self.status.setObjectName("status")
        bottom.addWidget(self.status, 1)
        outer.addLayout(bottom)
        root.addWidget(content, 1)
        self._actualizar_conteo_sistemas()
        self.setStyleSheet(STYLE)

    # Actualiza el indicador de cuántos sistemas tienen un archivo cargado.
    def _actualizar_conteo_sistemas(self):
        if not hasattr(self, "system_count_label") or not hasattr(
            self, "external_count_label"
        ):
            return
        internal_total = len(self.system_rows)
        internal_loaded = sum(
            bool(edit.file_path()) for edit in self.system_rows.values()
        )
        external_total = len(self.external_system_rows)
        external_loaded = sum(
            bool(edit.file_path()) for edit in self.external_system_rows.values()
        )
        self.system_count_label.setText(
            f"{internal_loaded} de {internal_total} cargados"
        )
        self.external_count_label.setText(
            f"{external_loaded} de {external_total} cargados"
        )

    @staticmethod
    def _default_output_path(filename_prefix: str) -> Path:
        documents = QStandardPaths.writableLocation(
            QStandardPaths.StandardLocation.DocumentsLocation
        )
        directory = Path(documents) if documents else Path.home() / "Documents"
        current_date = date.today().isoformat()
        return directory / f"{filename_prefix}_{current_date}.xlsx"

    def _guardar_excel_generado(
        self,
        temporal: str | Path,
        titulo: str,
        prefijo: str,
        nombre_sugerido: str | None = None,
    ) -> Path | None:
        """Pide la ruta solo cuando el XLSX temporal ya está generado."""
        temporal = Path(temporal)
        ruta_inicial = self._default_output_path(prefijo)
        if nombre_sugerido:
            ruta_inicial = ruta_inicial.with_name(nombre_sugerido)
        destino, _ = QFileDialog.getSaveFileName(
            self, titulo, str(ruta_inicial), "Libro de Excel (*.xlsx)"
        )
        if not destino:
            temporal.unlink(missing_ok=True)
            self.status.setText("Guardado cancelado; se descartó el Excel temporal.")
            return None

        destino_path = Path(destino)
        if destino_path.suffix.casefold() != ".xlsx":
            destino_path = destino_path.with_suffix(".xlsx")
        try:
            destino_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(temporal, destino_path)
        except Exception as exc:
            QMessageBox.critical(
                self,
                "Error al guardar el Excel",
                f"No se pudo guardar el archivo:\n{exc}\n\n"
                f"El Excel temporal se conserva en:\n{temporal}",
            )
            self.status.setText("No se pudo guardar el Excel generado.")
            return None

        temporal.unlink(missing_ok=True)
        return destino_path

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if not hasattr(self, "cards_splitter"):
            return
        compact = event.size().width() < 920
        self.sidebar.setFixedWidth(64 if compact else 220)
        self.sidebar_layout.setContentsMargins(
            6 if compact else 14, 22, 6 if compact else 14, 18
        )
        self.brand_logo_card.setVisible(not compact)
        self.brand_title.setVisible(compact)
        self.sidebar_footer.setVisible(not compact)
        self.nav_analysis_button.setText("◉" if compact else "◉   Analizar sistemas")
        self.nav_report_button.setText("▤" if compact else "▤   Generar Reporte")
        self.nav_telephony_button.setText("☎" if compact else "☎   Telefonía")
        self.nav_pdf_excel_button.setText("▧" if compact else PDF_EXCEL_NAV_TEXT)
        self.nav_pdf_excel_button.setMinimumHeight(42 if compact else 54)
        self.nav_analysis_button.setToolTip("Analizar sistemas")
        self.nav_report_button.setToolTip("Generar Reporte")
        self.nav_telephony_button.setToolTip("Telefonía")
        self.nav_pdf_excel_button.setToolTip("Conversor PDF a Excel")
        orientation = (
            Qt.Orientation.Horizontal
            if event.size().width() >= 1250
            else Qt.Orientation.Vertical
        )
        if self.cards_splitter.orientation() != orientation:
            self.cards_splitter.setOrientation(orientation)
            self._initial_split_applied = False
        if not self._initial_split_applied:
            QTimer.singleShot(0, self._apply_initial_split)

    def _show_view(self, index: int):
        self.views.setCurrentIndex(index)
        page_title = {
            0: "Analizador de Accesos",
            1: "Generar Reporte",
            2: "Analizador de Internos",
            3: "Conversor PDF a Excel",
        }.get(index, "Analizador de Accesos")
        self.page_title.setText(page_title)
        self.setWindowTitle(page_title)
        self.nav_analysis_button.setChecked(index == 0)
        self.nav_report_button.setChecked(index == 1)
        self.nav_telephony_button.setChecked(index == 2)
        self.nav_pdf_excel_button.setChecked(index == 3)

    def _apply_initial_split(self):
        if self.cards_splitter.orientation() == Qt.Orientation.Horizontal:
            width = self.cards_splitter.width()
            if width > 0:
                self.cards_splitter.setSizes([int(width * 0.55), int(width * 0.45)])
                self._initial_split_applied = True
        else:
            height = self.cards_splitter.height()
            if height > 0:
                self.cards_splitter.setSizes([height // 2, height // 2])
                self._initial_split_applied = True

    def _card(self, title: str) -> QFrame:
        frame = QFrame()
        frame.setObjectName("card")
        layout = QVBoxLayout(frame)
        layout.setContentsMargins(18, 14, 18, 16)
        layout.setSpacing(10)
        heading = QLabel(title)
        heading.setObjectName("sectionTitle")
        layout.addWidget(heading)
        return frame

    # Helpers visuales para filas de archivo, mensajes y selección de salida.
    def _file_row(
        self,
        parent: QWidget,
        label: str,
        file_filter: str,
        extensions: set[str],
        format_hint: str,
        configure=None,
        label_width: int | None = None,
        extra_action=None,
    ) -> DropPathEdit:
        row = QHBoxLayout()
        row.setContentsMargins(0, 1, 0, 1)
        identity = QVBoxLayout()
        identity.setSpacing(1)
        name = QLabel(label)
        name.setObjectName("fileName")
        if label_width:
            identity_widget = QWidget()
            identity_widget.setFixedWidth(label_width)
            identity_widget.setLayout(identity)
            name.setFixedWidth(label_width)
            name.setWordWrap(True)
            name.setMinimumHeight(self._wrapped_label_height(name, label_width))
        else:
            identity_widget = None
            name.setMinimumWidth(112)
        file_type = QLabel(format_hint)
        file_type.setObjectName("fileType")
        if label_width:
            file_type.setFixedWidth(label_width)
            file_type.setWordWrap(True)
            file_type.setMinimumHeight(
                self._wrapped_label_height(file_type, label_width)
            )
        identity.addWidget(name)
        identity.addWidget(file_type)
        edit = DropPathEdit(extensions)
        edit.setPlaceholderText("⇧   Arrastra aquí para cargar")
        edit.setToolTip(f"{label}: arrastra aquí o pulsa Buscar.")
        edit.dropRejected.connect(self._show_drop_error)
        path_hint = PathHintLabel()
        edit.fileChanged.connect(path_hint.set_path)
        file_stack = QWidget()
        file_stack.setObjectName("fileStack")
        file_stack_layout = QVBoxLayout(file_stack)
        file_stack_layout.setContentsMargins(0, 0, 0, 0)
        file_stack_layout.setSpacing(2)
        file_stack_layout.addWidget(edit)
        file_stack_layout.addWidget(path_hint)
        browse = QPushButton("Buscar")
        browse.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_DirOpenIcon))
        browse.setObjectName("browseButton")
        browse.setToolTip(f"Buscar archivo: {label}")
        browse.clicked.connect(
            lambda _=False: self._choose_file(edit, file_filter, label)
        )
        clear = QPushButton()
        clear.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_TrashIcon))
        clear.setObjectName("clearButton")
        clear.setToolTip(
            "Quitar SALAR actual e histórico"
            if extra_action
            else f"Quitar archivo: {label}"
        )
        if extra_action:
            clear.clicked.connect(lambda _=False: self._clear_salar_selection())
        else:
            clear.clicked.connect(lambda _=False: edit.set_file_path(""))
        if extra_action:
            history = QPushButton("+")
            history.setObjectName("historyButton")
            history.setToolTip("Cargar o revisar SALAR histórico")
            history.clicked.connect(lambda _checked=False: extra_action())
            self.salar_history_button = history
        if identity_widget:
            row.addWidget(identity_widget)
        else:
            row.addLayout(identity)
        row.addWidget(file_stack, 1)
        row.addWidget(browse)
        if extra_action:
            row.addWidget(history)
        row.addWidget(clear)
        if configure:
            gear = QPushButton("⚙")
            gear.setToolTip(f"Configurar {label}")
            gear.setObjectName("gearButton")
            gear.clicked.connect(lambda _checked=False, callback=configure: callback())
            row.addWidget(gear)
        parent.layout().addLayout(row)
        return edit

    @staticmethod
    def _wrapped_label_height(label: QLabel, width: int) -> int:
        metrics = QFontMetrics(label.font())
        # Dejar margen para el tamaño y el peso que aporta la hoja de estilos
        # después de construir los widgets.
        usable_width = max(1, int(width * 0.75))
        lines = 1
        current_line = ""
        for word in label.text().split():
            candidate = f"{current_line} {word}".strip()
            if current_line and metrics.horizontalAdvance(candidate) > usable_width:
                lines += 1
                current_line = word
            else:
                current_line = candidate
        return lines * metrics.lineSpacing() + 6

    def _report_file_box(self, title: str, required: bool):
        box = QFrame()
        box.setObjectName("reportInput")
        layout = QVBoxLayout(box)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(5)
        label = QLabel(title + (" *" if required else ""))
        label.setObjectName("fileName")
        layout.addWidget(label)

        row = QHBoxLayout()
        edit = DropPathEdit({".xlsx", ".xlsm"})
        edit.setPlaceholderText("⇧   Arrastra el Excel aquí")
        edit.setToolTip(f"{title}: arrastra aquí o pulsa Buscar.")
        edit.dropRejected.connect(self._show_drop_error)
        path_hint = PathHintLabel()
        edit.fileChanged.connect(path_hint.set_path)
        field_stack = QVBoxLayout()
        field_stack.setContentsMargins(0, 0, 0, 0)
        field_stack.setSpacing(2)
        field_stack.addWidget(edit)
        field_stack.addWidget(path_hint)
        row.addLayout(field_stack, 1)

        browse = QPushButton("Buscar")
        browse.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_DirOpenIcon))
        browse.setObjectName("browseButton")
        browse.clicked.connect(
            lambda _=False: self._choose_file(edit, MATRIX_FILTER, title)
        )
        row.addWidget(browse)
        clear = QPushButton()
        clear.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_TrashIcon))
        clear.setObjectName("clearButton")
        clear.setToolTip(f"Quitar archivo: {title}")
        clear.clicked.connect(lambda _=False: edit.set_file_path(""))
        row.addWidget(clear)
        layout.addLayout(row)
        return box, edit

    def _show_drop_error(self, message: str):
        self.status.setText(message)

    def _choose_file(self, edit: DropPathEdit, file_filter: str, label: str):
        path, _ = QFileDialog.getOpenFileName(
            self, f"Seleccionar {label}", "", file_filter
        )
        if path:
            if Path(path).suffix.casefold() not in edit.extensions:
                permitidas = ", ".join(sorted(edit.extensions))
                self._show_drop_error(
                    f"Formato inválido para {label}. Se esperaba: {permitidas}."
                )
                return
            edit.set_file_path(path)

    def _configure_salar_history(self):
        dialog = SalarHistoryDialog(self.salar_anterior_path, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.salar_anterior_path = dialog.file_edit.file_path()
            loaded = bool(self.salar_anterior_path)
            self.salar_history_button.setText("✓" if loaded else "+")
            self.salar_history_button.setProperty("historyLoaded", loaded)
            if loaded:
                self.salar_history_button.setToolTip(
                    "SALAR histórico cargado: "
                    f"{Path(self.salar_anterior_path).name}. Pulsa para revisar o cambiar."
                )
            else:
                self.salar_history_button.setToolTip(
                    "Cargar SALAR histórico (opcional)"
                )
            self.salar_history_button.style().unpolish(self.salar_history_button)
            self.salar_history_button.style().polish(self.salar_history_button)

    def _clear_salar_selection(self):
        self.salar_edit.set_file_path("")
        self.salar_anterior_path = ""
        self.salar_history_button.setText("+")
        self.salar_history_button.setProperty("historyLoaded", False)
        self.salar_history_button.setToolTip("Cargar SALAR histórico (opcional)")
        self.salar_history_button.style().unpolish(self.salar_history_button)
        self.salar_history_button.style().polish(self.salar_history_button)

    def _convert_pdf_to_excel(self):
        pdf_text = self.pdf_excel_input.file_path().strip()
        if not pdf_text or not Path(pdf_text).is_file():
            QMessageBox.warning(
                self,
                "Falta el PDF de ASFI",
                "Selecciona un archivo PDF válido del reporte de ASFI.",
            )
            return
        self.pdf_excel_suggested_name = Path(pdf_text).with_suffix(".xlsx").name
        self.convert_pdf_excel_button.setEnabled(False)
        self.status.setText("Generando el Excel con correos, nombres y roles…")
        self.pdf_excel_worker = PdfExcelWorker(Path(pdf_text))
        self.pdf_excel_worker.completed.connect(self._pdf_excel_done)
        self.pdf_excel_worker.failed.connect(self._pdf_excel_failed)
        self.pdf_excel_worker.start()

    def _pdf_excel_done(self, output: str, registros: int):
        self.convert_pdf_excel_button.setEnabled(True)
        destino_path = self._guardar_excel_generado(
            output,
            "Guardar Excel de roles ASFI",
            "roles_asfi",
            getattr(self, "pdf_excel_suggested_name", None),
        )
        if destino_path is None:
            return
        self._clear_pdf_excel_selection()
        self.status.setText(
            f"Excel ASFI guardado: {destino_path.name} · contiene {registros} registros"
        )
        QMessageBox.information(
            self,
            "Conversión completada",
            f"El Excel se generó y guardó correctamente:\n{destino_path}\n\n"
            f"Contiene {registros} registros. Cada rol ocupa una fila; "
            "los usuarios sin roles también se conservan.",
        )

    def _pdf_excel_failed(self, message: str):
        self.convert_pdf_excel_button.setEnabled(True)
        self.status.setText("No se pudo convertir el PDF de ASFI.")
        QMessageBox.critical(self, "Error al convertir el PDF", message)

    def _clear_pdf_excel_selection(self):
        self.pdf_excel_input.set_file_path("")
        self.status.setText("PDF quitado. Puedes seleccionar otro para convertirlo.")

    def _convert_sintesis_pdf_to_excel(self):
        pdf_text = self.sintesis_pdf_input.file_path().strip()
        if not pdf_text or not Path(pdf_text).is_file():
            QMessageBox.warning(
                self,
                "Falta el PDF de Síntesis",
                "Selecciona un archivo PDF válido del reporte de Síntesis.",
            )
            return

        self.sintesis_pdf_suggested_name = Path(pdf_text).with_suffix(".xlsx").name
        self.convert_sintesis_pdf_button.setEnabled(False)
        self.status.setText("Procesando el reporte PDF de Síntesis…")
        self.sintesis_excel_worker = SintesisExcelWorker(Path(pdf_text))
        self.sintesis_excel_worker.completed.connect(self._sintesis_excel_done)
        self.sintesis_excel_worker.failed.connect(self._sintesis_excel_failed)
        self.sintesis_excel_worker.start()

    def _sintesis_excel_done(self, output: str, registros: int):
        self.convert_sintesis_pdf_button.setEnabled(True)
        destino = self._guardar_excel_generado(
            output,
            "Guardar Excel de Síntesis",
            "sintesis",
            getattr(self, "sintesis_pdf_suggested_name", None),
        )
        if destino is None:
            return

        self._clear_sintesis_pdf_selection()
        self.status.setText(
            f"Excel de Síntesis guardado: {destino.name} · contiene {registros} registros"
        )
        QMessageBox.information(
            self,
            "Conversión de Síntesis completada",
            f"El Excel se generó y guardó correctamente:\n{destino}\n\n"
            f"Contiene {registros} registros.",
        )

    def _sintesis_excel_failed(self, message: str):
        self.convert_sintesis_pdf_button.setEnabled(True)
        self.status.setText("No se pudo convertir el PDF de Síntesis.")
        QMessageBox.critical(self, "Error al convertir el PDF de Síntesis", message)

    def _clear_sintesis_pdf_selection(self):
        self.sintesis_pdf_input.set_file_path("")
        self.status.setText("PDF de Síntesis quitado. Puedes seleccionar otro.")

    # Acciones de la ventana: validan entradas, inician workers y muestran resultados.
    def _generate_telephony_report(self):
        domain_text = self.telephony_domain_edit.file_path().strip()
        internal_text = self.telephony_internal_edit.file_path().strip()
        if not domain_text or not Path(domain_text).is_file():
            QMessageBox.warning(
                self, "Falta Dominio", "Selecciona un archivo TXT válido de Dominio."
            )
            return
        if not internal_text or not Path(internal_text).is_file():
            QMessageBox.warning(
                self,
                "Falta Internos (Isabel)",
                "Selecciona un archivo TXT válido de Internos (Isabel).",
            )
            return
        self.generate_telephony_button.setEnabled(False)
        self.status.setText("Analizando nombres y teléfonos…")
        output = _crear_ruta_temporal_excel("revision_telefonia_")
        self.telephony_worker = TelephonyWorker(
            Path(domain_text),
            Path(internal_text),
            output,
            self.source_configs["telefonia_dominio"],
            self.source_configs["telefonia_isabel"],
        )
        self.telephony_worker.completed.connect(self._telephony_done)
        self.telephony_worker.failed.connect(self._telephony_failed)
        self.telephony_worker.start()

    def _telephony_done(self, output: str, observations: int):
        self.generate_telephony_button.setEnabled(True)
        destino_path = self._guardar_excel_generado(
            output, "Guardar análisis de telefonía", "revision_telefonia"
        )
        if destino_path is None:
            return
        self._clear_telephony_selection()
        self.status.setText(
            f"Análisis de telefonía guardado: {destino_path.name} · "
            f"{observations} observaciones"
        )
        QMessageBox.information(
            self,
            "Análisis de telefonía generado",
            f"El reporte se guardó en:\n{destino_path}\n\n"
            f"Observaciones generadas: {observations}\n"
            "El Excel contiene las categorias detectadas en el analisis.",
        )

    def _telephony_failed(self, message: str):
        self.generate_telephony_button.setEnabled(True)
        self.status.setText("No se pudo generar el análisis de telefonía.")
        QMessageBox.critical(self, "Error en el análisis de telefonía", message)

    def _clear_telephony_selection(self):
        self.telephony_domain_edit.set_file_path("")
        self.telephony_internal_edit.set_file_path("")
        self.status.setText("Archivos de telefonía quitados.")

    def _generate_report(self):
        analysis_text = self.report_analysis_edit.file_path().strip()
        if not analysis_text or not Path(analysis_text).is_file():
            QMessageBox.warning(
                self,
                "Falta la revisión de accesos",
                "Carga un Excel válido de revisión de accesos.",
            )
            return
        false_positive_text = self.false_positives_edit.file_path().strip()
        if false_positive_text and not Path(false_positive_text).is_file():
            QMessageBox.warning(
                self,
                "Falsos positivos no válido",
                "Selecciona un Excel válido o quita el archivo opcional.",
            )
            return
        self.generate_report_button.setEnabled(False)
        self.status.setText("Generando informe…")
        false_positive_file = Path(false_positive_text) if false_positive_text else None
        output = _crear_ruta_temporal_excel("reporte_revision_")
        self.report_worker = ReportWorker(
            Path(analysis_text), output, false_positive_file
        )
        self.report_worker.completed.connect(self._report_done)
        self.report_worker.failed.connect(self._report_failed)
        self.report_worker.start()

    def _report_done(self, output: str):
        self.generate_report_button.setEnabled(True)
        destino_path = self._guardar_excel_generado(
            output, "Guardar informe", "reporte_revision"
        )
        if destino_path is None:
            return
        self._clear_report_selection()
        self.status.setText(f"Informe guardado: {destino_path.name}")
        QMessageBox.information(
            self, "Informe generado", f"El informe se guardó en:\n{destino_path}"
        )

    def _report_failed(self, message: str):
        self.generate_report_button.setEnabled(True)
        self.status.setText("No se pudo generar el informe.")
        QMessageBox.critical(self, "Error al generar el informe", message)

    # Abre los diálogos de configuración y aplica sus valores a la sesión.
    def _configure(self, name: str):
        dialog = ConfigDialog(self.configs[name], self)
        if dialog.exec():
            self.configs[name] = dialog.config
            self._save_config(dialog.config)
            self.status.setText(f"Configuración de {name} guardada.")

    def _configure_source(self, key: str, title: str):
        dialog = SourceConfigDialog(title, self.source_configs[key], self)
        if dialog.exec():
            self.source_configs[key] = dialog.config
            self._save_source_config(key, dialog.config)
            self.status.setText(f"Configuración de {title} guardada.")

    # Valida archivos seleccionados y lanza la revisión de accesos en segundo plano.
    def _analyze(self):
        # La fuente de personas es fija: SALAR; Dominio sirve para conciliar identidades.
        source = "SALAR"
        matrix_path = self.matrix_edit.file_path()
        domain_path = self.domain_edit.file_path()
        matrix = Path(matrix_path) if matrix_path else None
        domain = Path(domain_path) if domain_path else None
        if not matrix or not matrix.is_file():
            QMessageBox.warning(
                self, "Falta la matriz", "Selecciona un archivo Excel de matriz válido."
            )
            return
        salar_path = self.salar_edit.file_path()
        people_file = Path(salar_path) if salar_path else None
        if not domain or not domain.is_file():
            QMessageBox.warning(
                self,
                "Falta Dominio",
                "Selecciona un archivo TXT de Dominio válido para identificar usuarios externos.",
            )
            return
        if not people_file or not people_file.is_file():
            QMessageBox.warning(
                self,
                f"Falta {source}",
                f"Selecciona un archivo TXT de {source} válido.",
            )
            return
        selected = [
            (name, path) for name, path in self.system_rows.items() if path.file_path()
        ]
        asfi_text = self.asfi_edit.file_path().strip()
        asfi_file = Path(asfi_text) if asfi_text else None
        infocred_text = self.infocred_edit.file_path().strip()
        infocred_file = Path(infocred_text) if infocred_text else None
        if not selected and not asfi_file and not infocred_file:
            QMessageBox.warning(
                self,
                "Sin sistemas",
                "Carga el TXT de al menos un sistema para analizar.",
            )
            return
        missing = [
            name
            for name, path in selected
            if not path.file_path() or not Path(path.file_path()).is_file()
        ]
        if missing:
            QMessageBox.warning(
                self,
                "Falta seleccionar TXT",
                "Verifica el archivo TXT de: " + ", ".join(missing),
            )
            return
        if asfi_file and not asfi_file.is_file():
            QMessageBox.warning(
                self,
                "TXT ASFI no válido",
                "Selecciona un archivo ASFI válido o quita el archivo cargado.",
            )
            return
        if infocred_file and not infocred_file.is_file():
            QMessageBox.warning(
                self,
                "TXT INFOCRED no válido",
                "Selecciona un archivo INFOCRED válido o quita el archivo cargado.",
            )
            return
        files = {name: Path(path.file_path()) for name, path in selected}
        configs = tuple(self.configs[name] for name, _ in selected)
        temporal_text = self.temporales_edit.file_path().strip()
        temporales = Path(temporal_text) if temporal_text else None
        if temporales and not temporales.is_file():
            QMessageBox.warning(
                self,
                "Excel temporal no válido",
                "Selecciona un archivo Excel temporal válido o deja el campo vacío.",
            )
            return
        salar_anterior_text = self.salar_anterior_path.strip()
        salar_anterior = Path(salar_anterior_text) if salar_anterior_text else None
        if salar_anterior and not salar_anterior.is_file():
            QMessageBox.warning(
                self,
                "SALAR anterior no válido",
                "Selecciona un TXT SALAR anterior válido o deja el campo vacío.",
            )
            return
        output = _crear_ruta_temporal_excel("revision_accesos_")
        self.analyze_salar_button.setEnabled(False)
        self.status.setText(f"Analizando con {source}…")
        person_config = self.source_configs[source.casefold()]
        self.worker = AnalysisWorker(
            matrix,
            people_file,
            source,
            files,
            output,
            configs,
            temporales,
            self.source_configs["matriz"],
            person_config,
            self.source_configs["temporales"],
            domain,
            self.source_configs["dominio"],
            asfi_file,
            self.source_configs["asfi"],
            salar_anterior,
            infocred_file,
            self.source_configs["infocred"],
        )
        self.worker.completed.connect(self._analysis_done)
        self.worker.failed.connect(self._analysis_failed)
        self.worker.start()

    def _analysis_done(self, output: str, observations: int, errors: int):
        self.analyze_salar_button.setEnabled(True)
        destino_path = self._guardar_excel_generado(
            output, "Guardar revisión de accesos", "revision_accesos"
        )
        if destino_path is None:
            return
        self._clear_selection()
        self.status.setText(
            f"Reporte guardado: {destino_path.name} · {observations} observaciones · {errors} errores"
        )
        QMessageBox.information(
            self,
            "Análisis completado",
            f"El reporte se guardó en:\n{destino_path}\n\n"
            f"Observaciones: {observations}\nErrores de carga: {errors}",
        )

    def _analysis_failed(self, message: str):
        self.analyze_salar_button.setEnabled(True)
        self.status.setText("No se pudo completar el análisis.")
        QMessageBox.critical(self, "Error durante el análisis", message)

    def _clear_selection(self):
        self.matrix_edit.set_file_path("")
        self.domain_edit.set_file_path("")
        self.salar_edit.set_file_path("")
        self.salar_anterior_path = ""
        self.salar_history_button.setText("+")
        self.salar_history_button.setProperty("historyLoaded", False)
        self.salar_history_button.setToolTip("Cargar SALAR histórico (opcional)")
        self.salar_history_button.style().unpolish(self.salar_history_button)
        self.salar_history_button.style().polish(self.salar_history_button)
        self.temporales_edit.set_file_path("")
        self.asfi_edit.set_file_path("")
        self.infocred_edit.set_file_path("")
        for path in self.system_rows.values():
            path.set_file_path("")
        self.status.setText("Archivos quitados. La configuración se conserva.")

    def _clear_report_selection(self):
        self.report_analysis_edit.set_file_path("")
        self.false_positives_edit.set_file_path("")
        self.status.setText("Archivos del informe quitados.")


STYLE = """
QMainWindow, QWidget { background: #edf2f7; color: #243247; font-family: 'Segoe UI'; font-size: 10pt; }
QLabel { background: transparent; }
QWidget#content { background: #edf2f7; }
QFrame#sidebar { background: #10275c; border: 0; }
QFrame#sidebar QLabel { background: transparent; }
QFrame#brandLogoCard { background: #ffffff; border: 0; border-radius: 9px; }
QLabel#brandLogo { background: transparent; }
QLabel#brandTitle { color: white; font-size: 25pt; font-weight: 800; padding: 2px; }
QLabel#navButton { color: white; background: #24417f; border: 0; border-radius: 8px; font-weight: 700; }
QPushButton#navButton { color: white; background: transparent; border: 0; border-radius: 8px; padding: 8px 10px; text-align: left; font-weight: 700; }
QPushButton#navButton:hover, QPushButton#navButton:checked { background: #24417f; }
QLabel#sidebarFooter { color: #c4d2f2; font-size: 8pt; }
QLabel#eyebrow { color: #4f63ad; font-size: 8pt; font-weight: 700; letter-spacing: 1px; }
QLabel#title { color: #14253d; font-size: 24pt; font-weight: 700; padding: 1px; }
QLabel#subtitle, QLabel#hint { color: #65748a; }
QLabel#subtitle { font-size: 10pt; }
QLabel#sectionTitle { color: #263b58; font-size: 12pt; font-weight: 700; }
QLabel#versionBadge, QLabel#countBadge { color: #315287; background: #e5edfb; border: 1px solid #d4e0f4; border-radius: 10px; padding: 5px 10px; font-size: 8pt; font-weight: 700; }
QFrame#card, QFrame#actionCard { background: #ffffff; border: 1px solid #d9e2ec; border-radius: 12px; }
QFrame#systemsCard { background: #ffffff; border: 1px solid #d9e2ec; border-radius: 12px; }
QScrollArea#contentScroll, QScrollArea#systemRowsScroll { border: 0; background: transparent; }
QWidget#contentPage, QWidget#systemRowsContent { background: transparent; }
QWidget#fileStack { background: transparent; border: none; }
QSplitter::handle { background: #e3eaf3; }
QSplitter::handle:horizontal { width: 8px; }
QSplitter::handle:vertical { height: 8px; }
QLineEdit { background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 7px; padding: 8px 10px; }
QLineEdit:focus { border: 1px solid #3976c5; }
QLineEdit#filePicker { background: #f7faff; border: 1px dashed #9aaec8; color: #6c7e95; padding: 7px 10px; }
QLineEdit#filePicker:focus { border: 1px dashed #3976c5; }
QLineEdit#filePicker[fileLoaded="true"] { background: #e9f8ee; border: 1px solid #9fd4b2; color: #13834c; font-size: 15pt; font-weight: 800; }
QLabel#fileName { color: #263b58; font-weight: 700; }
QLabel#fileType { color: #718198; font-size: 8pt; }
QLabel#pathHint { color: #91a0b3; font-size: 8pt; padding-left: 5px; }
QPushButton { background: #ffffff; border: 1px solid #cbd5e1; border-radius: 7px; padding: 8px 12px; }
QPushButton:hover { background: #edf4ff; border-color: #7aa4dc; }
QPushButton#secondaryButton { color: #29436b; font-weight: 600; }
QPushButton#browseButton { min-height: 36px; padding: 6px 10px; color: #29436b; font-weight: 600; }
QPushButton#clearButton, QPushButton#gearButton { min-width: 36px; max-width: 36px; min-height: 36px; padding: 2px; font-size: 13pt; }
QPushButton#historyButton { min-width: 36px; max-width: 36px; min-height: 36px; padding: 2px; font-size: 18pt; font-weight: 600; color: #264f84; }
QPushButton#historyButton[historyLoaded="true"] { color: #24734a; border-color: #80bd9b; background: #eef8f1; }
QPushButton#clearButton { color: #a33c46; }
QPushButton#clearButton:hover { background: #fff1f1; border-color: #e6a5aa; }
QPushButton#gearButton { color: #264f84; }
QPushButton#analyzeButton { min-width: 218px; color: #ffffff; background-color: #2454b8; border: 1px solid #1d469e; border-radius: 9px; padding: 11px 18px; font-weight: 700; }
QPushButton#analyzeButton:hover { color: #ffffff; background-color: #2d60c6; border-color: #2454b8; }
QPushButton#analyzeButton:pressed { background-color: #1d469e; border-color: #183b86; }
QPushButton#analyzeButton:disabled { color: #ffffff; background-color: #91a9d7; border-color: #91a9d7; }
QCheckBox { spacing: 8px; font-weight: 600; }
QCheckBox::indicator { width: 17px; height: 17px; }
QLabel#status { color: #52647b; padding: 2px 4px; }
"""


def main() -> int:
    if sys.platform == "win32":
        import ctypes

        set_app_id = ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID
        set_app_id.argtypes = [ctypes.c_wchar_p]
        set_app_id.restype = ctypes.c_long
        result = set_app_id(APP_USER_MODEL_ID)
        if result != 0:
            raise OSError(
                f"No se pudo configurar el identificador de Windows: {result}"
            )
    app = QApplication(sys.argv)
    app.setOrganizationName("RevisionAccesos")
    app.setApplicationName("Analizador de Accesos")
    app.setApplicationDisplayName("Analizador de Accesos")
    icon_path = recurso_aplicacion("assets/icono.ico")
    if icon_path.is_file():
        app.setWindowIcon(QIcon(str(icon_path)))
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
